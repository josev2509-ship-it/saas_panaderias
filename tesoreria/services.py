import csv
import hashlib
import io
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone

from auditoria.models import EventoAuditoria
from auditoria.services import registrar_evento
from contabilidad.domain_events import ConciliacionCompletada
from core.application.event_bus import event_bus

from .models import ConciliacionBancaria, ImportacionExtractoBancario, LineaConciliacion, LineaExtractoBancario, MovimientoTesoreria


DEFAULT_MAPPING = {"fecha": "fecha", "descripcion": "descripcion", "referencia": "referencia", "debito": "debito", "credito": "credito", "saldo": "saldo"}


def _decimal(value):
    try:
        return Decimal(str(value or "0").replace(",", "").strip())
    except InvalidOperation as exc:
        raise ValidationError("El extracto contiene un importe inválido.") from exc


def _date(value):
    if hasattr(value, "year") and not isinstance(value, str):
        return value.date() if hasattr(value, "date") else value
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(str(value).strip(), formato).date()
        except ValueError:
            pass
    raise ValidationError("El extracto contiene una fecha inválida.")


def _normalizar(rows, mapeo=None):
    mapeo = {**DEFAULT_MAPPING, **(mapeo or {})}
    filas = []
    for numero, row in enumerate(rows, 2):
        try:
            fecha = _date(row[mapeo["fecha"]])
            debito, credito = _decimal(row.get(mapeo["debito"])), _decimal(row.get(mapeo["credito"]))
            if debito and credito:
                raise ValidationError(f"Fila {numero}: no puede contener débito y crédito.")
            referencia = str(row.get(mapeo["referencia"], "") or "").strip()[:100]
            monto = credito - debito
            huella = hashlib.sha256(f"{fecha}|{referencia}|{monto}".encode()).hexdigest()
            filas.append({"linea": numero, "fecha": fecha, "descripcion": str(row.get(mapeo["descripcion"], "") or "").strip()[:255], "referencia": referencia, "monto": monto, "saldo": _decimal(row.get(mapeo["saldo"])), "huella": huella})
        except KeyError as exc:
            raise ValidationError(f"Fila {numero}: falta la columna {exc.args[0]}.") from exc
    return filas


def previsualizar_csv(contenido, *, mapeo=None):
    text = contenido.decode("utf-8-sig") if isinstance(contenido, bytes) else contenido
    return _normalizar(csv.DictReader(io.StringIO(text)), mapeo)


def previsualizar_xlsx(contenido, *, mapeo=None):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValidationError("La importación XLSX no está disponible.") from exc
    archivo = io.BytesIO(contenido) if isinstance(contenido, bytes) else contenido
    hoja = load_workbook(archivo, read_only=True, data_only=True).active
    valores = hoja.iter_rows(values_only=True)
    try:
        encabezados = [str(x or "").strip() for x in next(valores)]
    except StopIteration as exc:
        raise ValidationError("El archivo XLSX está vacío.") from exc
    return _normalizar((dict(zip(encabezados, row)) for row in valores), mapeo)


@transaction.atomic
def importar_extracto(*, context, cuenta, contenido, nombre_archivo, mapeo=None, confirmar=False):
    if cuenta.empresa_id != context.empresa.pk:
        raise ValidationError("La cuenta no pertenece a la empresa.")
    raw = contenido if isinstance(contenido, bytes) else contenido.encode("utf-8")
    formato = nombre_archivo.rsplit(".", 1)[-1].upper()
    if formato not in {"CSV", "XLSX"}:
        raise ValidationError("Solo se admiten extractos CSV o XLSX.")
    filas = previsualizar_csv(raw, mapeo=mapeo) if formato == "CSV" else previsualizar_xlsx(raw, mapeo=mapeo)
    if not confirmar:
        return filas
    huella = hashlib.sha256(raw).hexdigest()
    importacion, creada = ImportacionExtractoBancario.objects.get_or_create(empresa=context.empresa, cuenta=cuenta, huella=huella, defaults={"nombre_archivo": nombre_archivo[:180], "formato": formato, "estado": "CONFIRMADA", "mapeo": mapeo or DEFAULT_MAPPING, "confirmado_por": context.usuario, "confirmado_en": timezone.now()})
    if creada:
        LineaExtractoBancario.objects.bulk_create([LineaExtractoBancario(importacion=importacion, numero=f["linea"], fecha=f["fecha"], descripcion=f["descripcion"], referencia=f["referencia"], monto=f["monto"], saldo=f["saldo"], huella=f["huella"]) for f in filas])
        registrar_evento(empresa=context.empresa, usuario=context.usuario, request=context.request, objeto=importacion, modulo="tesoreria", accion=EventoAuditoria.Accion.OTRO, descripcion="Extracto bancario importado.", datos_nuevos={"formato": formato, "filas": len(filas), "correlation_id": context.identificador_solicitud})
    return importacion


def importar_csv(*, empresa, cuenta, contenido, usuario=None, mapeo=None, confirmar=False):
    from core.application.operation_context import OperationContext
    return importar_extracto(context=OperationContext(empresa=empresa, usuario=usuario), cuenta=cuenta, contenido=contenido, nombre_archivo="extracto.csv", mapeo=mapeo, confirmar=confirmar)


@transaction.atomic
def sugerir_coincidencias(*, context, importacion, tolerancia_dias=3):
    importacion = ImportacionExtractoBancario.objects.get(pk=importacion.pk, empresa=context.empresa)
    sugerencias = []
    lineas = list(importacion.lineas.filter(estado="PENDIENTE"))
    if not lineas:
        return sugerencias
    desde = min(x.fecha for x in lineas) - timedelta(days=tolerancia_dias)
    hasta = max(x.fecha for x in lineas) + timedelta(days=tolerancia_dias)
    movimientos = list(MovimientoTesoreria.objects.filter(empresa=context.empresa, cuenta=importacion.cuenta, conciliado=False, fecha__range=(desde, hasta)).order_by("fecha", "pk"))
    for linea in lineas:
        candidatos = [x for x in movimientos if x.monto == abs(linea.monto) and abs((x.fecha-linea.fecha).days) <= tolerancia_dias]
        exacto = next((x for x in candidatos if linea.referencia and x.referencia == linea.referencia), None)
        elegido = exacto or (candidatos[0] if candidatos else None)
        if elegido:
            sugerencias.append({"linea_id": linea.pk, "movimiento_id": elegido.pk, "tipo": "REFERENCIA" if exacto else "MONTO_FECHA"})
    return sugerencias

def detectar_ambiguedades(*,context,importacion,tolerancia_dias=3):
    importacion=ImportacionExtractoBancario.objects.get(pk=importacion.pk,empresa=context.empresa);resultado=[]
    for linea in importacion.lineas.filter(estado="PENDIENTE"):
        candidatos=list(MovimientoTesoreria.objects.filter(empresa=context.empresa,cuenta=importacion.cuenta,conciliado=False,fecha__range=(linea.fecha-timedelta(days=tolerancia_dias),linea.fecha+timedelta(days=tolerancia_dias)),monto=abs(linea.monto)).order_by("fecha","pk"))
        if len(candidatos)>1:resultado.append({"linea_id":linea.pk,"candidatos":[{"movimiento_id":x.pk,"fecha":str(x.fecha),"monto":str(x.monto),"referencia":x.referencia} for x in candidatos],"tipo":"AMBIGUO"})
    return resultado

@transaction.atomic
def seleccionar_coincidencia_manual(*,context,conciliacion,linea_id,movimiento_id,motivo):
    if not motivo:raise ValidationError("El motivo de selección manual es obligatorio.")
    linea=LineaExtractoBancario.objects.select_for_update().get(pk=linea_id,importacion__empresa=context.empresa,importacion__cuenta=conciliacion.cuenta,estado="PENDIENTE");movimiento=MovimientoTesoreria.objects.select_for_update().get(pk=movimiento_id,empresa=context.empresa,cuenta=conciliacion.cuenta,conciliado=False)
    candidatos=detectar_ambiguedades(context=context,importacion=linea.importacion)
    permitidos={x["movimiento_id"] for item in candidatos if item["linea_id"]==linea.pk for x in item["candidatos"]}
    if movimiento.pk not in permitidos:raise ValidationError("El movimiento no es candidato válido para la línea.")
    resultado=conciliar_linea(context=context,conciliacion=conciliacion,linea_id=linea.pk,movimiento_id=movimiento.pk,tipo="MANUAL")
    registrar_evento(empresa=context.empresa,usuario=context.usuario,request=context.request,objeto=linea,modulo="tesoreria",accion=EventoAuditoria.Accion.OTRO,descripcion="Coincidencia bancaria ambigua resuelta manualmente.",datos_nuevos={"movimiento_id":movimiento.pk,"motivo":motivo,"correlation_id":context.identificador_solicitud});return resultado

@transaction.atomic
def reconciliar_linea(*,context,conciliacion,linea_id,movimiento_id,motivo):
    linea=LineaExtractoBancario.objects.get(pk=linea_id,importacion__empresa=context.empresa)
    if linea.estado=="CONCILIADA" and linea.movimiento_id==movimiento_id:return linea
    if linea.estado!="PENDIENTE":raise ValidationError("La línea no está disponible para reconciliar.")
    resultado=conciliar_linea(context=context,conciliacion=conciliacion,linea_id=linea_id,movimiento_id=movimiento_id,tipo="RECONCILIACION")
    registrar_evento(empresa=context.empresa,usuario=context.usuario,request=context.request,objeto=resultado,modulo="tesoreria",accion=EventoAuditoria.Accion.OTRO,descripcion="Línea bancaria reconciliada.",datos_nuevos={"motivo":motivo});return resultado


@transaction.atomic
def conciliar_linea(*, context, conciliacion, linea_id, movimiento_id, tipo="MANUAL"):
    conciliacion = ConciliacionBancaria.objects.select_for_update().get(pk=conciliacion.pk, empresa=context.empresa, estado="BORRADOR")
    linea = LineaExtractoBancario.objects.select_for_update().get(pk=linea_id, importacion__empresa=context.empresa, importacion__cuenta=conciliacion.cuenta, estado="PENDIENTE")
    movimiento = MovimientoTesoreria.objects.select_for_update().get(pk=movimiento_id, empresa=context.empresa, cuenta=conciliacion.cuenta, conciliado=False)
    if abs(linea.monto) != movimiento.monto:
        raise ValidationError("Los importes no coinciden.")
    LineaConciliacion.objects.create(conciliacion=conciliacion, movimiento=movimiento, descripcion=linea.descripcion, monto_banco=linea.monto, monto_libros=movimiento.monto if movimiento.tipo == "INGRESO" else -movimiento.monto, coincide=True)
    linea.movimiento, linea.estado, linea.tipo_coincidencia = movimiento, "CONCILIADA", tipo
    linea.save(update_fields=["movimiento", "estado", "tipo_coincidencia"])
    movimiento.conciliado = True
    movimiento.save(update_fields=["conciliado"])
    return linea


@transaction.atomic
def desconciliar_linea(*, context, linea_id, motivo):
    if not motivo:
        raise ValidationError("El motivo es obligatorio.")
    linea = LineaExtractoBancario.objects.select_for_update().get(pk=linea_id, importacion__empresa=context.empresa, estado="CONCILIADA")
    MovimientoTesoreria.objects.filter(pk=linea.movimiento_id, empresa=context.empresa).update(conciliado=False)
    LineaConciliacion.objects.filter(movimiento_id=linea.movimiento_id, conciliacion__empresa=context.empresa, conciliacion__estado="BORRADOR").delete()
    linea.movimiento = None
    linea.estado = "PENDIENTE"
    linea.tipo_coincidencia = ""
    linea.save(update_fields=["movimiento", "estado", "tipo_coincidencia"])
    return linea


@transaction.atomic
def completar_conciliacion(*, context, conciliacion):
    conciliacion = ConciliacionBancaria.objects.select_for_update().get(pk=conciliacion.pk, empresa=context.empresa)
    if conciliacion.estado == "CERRADA":
        return conciliacion
    if conciliacion.saldo_banco != conciliacion.saldo_libros or conciliacion.lineas.filter(coincide=False).exists():
        raise ValidationError("La conciliación contiene diferencias pendientes.")
    conciliacion.estado = "CERRADA"
    conciliacion.save(update_fields=["estado"])
    registrar_evento(empresa=context.empresa, usuario=context.usuario, request=context.request, objeto=conciliacion, modulo="tesoreria", accion=EventoAuditoria.Accion.CAMBIAR_ESTADO, descripcion="Conciliación bancaria cerrada.", datos_nuevos={"estado": "CERRADA", "correlation_id": context.identificador_solicitud})
    event_bus.publish(ConciliacionCompletada(empresa_id=context.empresa.pk, usuario_id=getattr(context.usuario, "pk", None), agregado_tipo="tesoreria.conciliacionbancaria", agregado_id=str(conciliacion.pk), referencia=str(conciliacion.pk), clave_idempotente=f"conciliacion-completada:{conciliacion.pk}", payload={"schema_version": 1, "empresa_id": context.empresa.pk, "conciliacion_id": conciliacion.pk}))
    return conciliacion
