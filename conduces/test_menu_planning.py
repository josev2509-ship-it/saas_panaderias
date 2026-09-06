from datetime import date, timedelta
from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from auditoria.models import EventoAuditoria

from conduces.menu_planning import (
    activar_calendario,
    aplicar_excepcion,
    contar_docencia_regular,
    materializar_programacion,
    previsualizar_asignacion,
    semana_ciclo,
    semana_ciclo_escolar_regular,
    version_vigente,
)
from conduces.models import (
    AsignacionProgramaCentro,
    CalendarioEscolar,
    CentroEducativo,
    DiaCalendarioEscolar,
    Empresa,
    ItemCicloMenu,
    ProgramaMenu,
    ProgramacionMenuEscolar,
    VersionProgramaMenu,
)


class MenuPlanningTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("planificador", password="x")
        self.empresa = Empresa.objects.create(usuario=self.user, nombre="Panaderia A", modulo_menu=True)
        self.user.user_permissions.add(Permission.objects.get(codename="change_programamenu"))
        self.centro = CentroEducativo.objects.create(
            empresa=self.empresa, codigo="A-1", nombre="Centro A", matricula=100
        )
        self.calendario = CalendarioEscolar.objects.create(
            empresa=self.empresa,
            nombre="2026-2027",
            anio_inicio=2026,
            anio_fin=2027,
            inicio_docencia=date(2026, 8, 24),
            fin_docencia=date(2026, 9, 6),
            dias_docencia_oficiales=9,
            estado=CalendarioEscolar.Estado.ACTIVO,
        )
        fecha = self.calendario.inicio_docencia
        while fecha <= self.calendario.fin_docencia:
            DiaCalendarioEscolar.objects.create(
                calendario=self.calendario,
                fecha=fecha,
                clasificacion=(
                    DiaCalendarioEscolar.Clasificacion.DOCENCIA
                    if fecha.weekday() < 5
                    else DiaCalendarioEscolar.Clasificacion.NO_LECTIVO
                ),
            )
            fecha += timedelta(days=1)
        self.programa = ProgramaMenu.objects.create(
            empresa=self.empresa, codigo="URB-V2", nombre="PAE Urbano V2", modalidad=ProgramaMenu.Modalidad.REGULAR
        )
        self.version = VersionProgramaMenu.objects.create(
            programa=self.programa,
            nombre="V1",
            vigente_desde=date(2026, 8, 24),
            fecha_uso_desde=date(2026, 8, 24),
            semanas_ciclo=5,
            fecha_ancla_ciclo=date(2026, 8, 24),
            semana_inicial=1,
            estado=VersionProgramaMenu.Estado.ACTIVA,
        )
        productos = ("Pan", "Galleta", "Muffin", "Pan integral", "Bizcocho")
        for semana in range(1, 6):
            for dia in range(5):
                ItemCicloMenu.objects.create(
                    version=self.version,
                    semana=semana,
                    dia_semana=dia,
                    producto=f"{productos[dia]} S{semana}",
                )
        self.asignacion = AsignacionProgramaCentro.objects.create(
            centro=self.centro,
            programa=self.programa,
            modalidad=ProgramaMenu.Modalidad.REGULAR,
            dias_entrega=[0, 1, 2, 3, 4],
            vigente_desde=date(2026, 8, 24),
            creado_por=self.user,
        )

    def test_conteo_regular_excluye_sabado_y_domingo(self):
        sabado = self.calendario.dias.get(fecha=date(2026, 8, 29))
        sabado.clasificacion = DiaCalendarioEscolar.Clasificacion.DOCENCIA
        sabado.save()
        self.assertEqual(contar_docencia_regular(self.calendario), 10)

    def test_calendario_no_activa_si_conteo_no_coincide(self):
        self.calendario.estado = CalendarioEscolar.Estado.EN_REVISION
        self.calendario.dias_docencia_oficiales = 190
        self.calendario.save()
        with self.assertRaisesMessage(ValidationError, "Dias oficiales: 190"):
            activar_calendario(self.calendario, usuario=self.user)

    def test_previsualizacion_2026_2027_aplica_dia_oficial_adp(self):
        self.calendario.delete()
        self.client.force_login(self.user)
        response = self.client.post(reverse("crear_calendario_escolar_planificacion"), {
            "nombre": "Oficial 2026-2027",
            "anio_inicio": 2026,
            "anio_fin": 2027,
            "inicio_docencia": "2026-08-24",
            "fin_docencia": "2027-06-18",
            "dias_docencia_oficiales": 190,
        })
        self.assertRedirects(response, reverse("planificacion_menu_escolar"))
        calendario = CalendarioEscolar.objects.get(nombre="Oficial 2026-2027")
        adp = calendario.dias.get(fecha=date(2027, 4, 13))
        self.assertEqual(adp.clasificacion, DiaCalendarioEscolar.Clasificacion.NO_LECTIVO)
        self.assertEqual(adp.motivo, "Día de la ADP")
        self.assertEqual(adp.origen, "CATALOGO_OFICIAL")

    def test_carga_excel_oficial_aplica_excepciones_con_trazabilidad(self):
        self.calendario.estado = CalendarioEscolar.Estado.EN_REVISION
        self.calendario.save()
        wb = __import__("openpyxl").Workbook()
        ws = wb.active
        ws.append(["fecha", "clasificacion", "motivo"])
        ws.append(["2026-08-25", "FERIADO", "Feriado oficial"])
        contenido = BytesIO(); wb.save(contenido)
        archivo = SimpleUploadedFile(
            "fechas.xlsx", contenido.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("cargar_fechas_calendario_oficial", args=[self.calendario.pk]),
            {"archivo": archivo},
        )
        self.assertEqual(response.status_code, 302)
        dia = self.calendario.dias.get(fecha=date(2026, 8, 25))
        self.assertEqual(dia.clasificacion, DiaCalendarioEscolar.Clasificacion.FERIADO)
        self.assertEqual(dia.motivo, "Feriado oficial")
        self.assertEqual(dia.origen, "CARGA_OFICIAL_EXCEL")
        self.assertEqual(dia.ajustado_por, self.user)

    def test_revision_muestra_resumen_busqueda_filtro_y_diferencia(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("planificacion_menu_escolar"), {
            "calendario": self.calendario.pk,
            "fecha": "2026-08-24",
            "clasificacion": "DOCENCIA",
        })
        self.assertContains(response, "Días oficiales")
        self.assertContains(response, "Días calculados")
        self.assertContains(response, "Diferencia")
        self.assertContains(response, "24/08/2026")
        self.assertNotContains(response, "25/08/2026")

    def test_ajuste_manual_excepcional_exige_motivo(self):
        self.calendario.estado = CalendarioEscolar.Estado.EN_REVISION
        self.calendario.save()
        dia = self.calendario.dias.get(fecha=date(2026, 8, 24))
        self.client.force_login(self.user)
        response = self.client.post(reverse("clasificar_dia_calendario", args=[dia.pk]), {
            "clasificacion": "NO_LECTIVO",
            "motivo": "",
        })
        self.assertEqual(response.status_code, 302)
        dia.refresh_from_db()
        self.assertEqual(dia.clasificacion, DiaCalendarioEscolar.Clasificacion.DOCENCIA)

    def test_correccion_manual_con_motivo_es_auditable(self):
        self.calendario.estado = CalendarioEscolar.Estado.EN_REVISION
        self.calendario.save()
        dia = self.calendario.dias.get(fecha=date(2026, 8, 24))
        self.client.force_login(self.user)
        response = self.client.post(reverse("clasificar_dia_calendario", args=[dia.pk]), {
            "clasificacion": "SUSPENSION",
            "motivo": "Suspension oficial extraordinaria",
        })
        self.assertEqual(response.status_code, 302)
        dia.refresh_from_db()
        self.assertEqual(dia.clasificacion, "SUSPENSION")
        self.assertEqual(dia.ajustado_por, self.user)
        evento = EventoAuditoria.objects.get(object_id=dia.pk, modulo="planificacion_menu")
        self.assertEqual(evento.usuario, self.user)
        self.assertEqual(evento.datos_nuevos["clasificacion"], "SUSPENSION")

    def test_feriado_no_desplaza_producto_del_miercoles(self):
        martes = self.calendario.dias.get(fecha=date(2026, 8, 25))
        martes.clasificacion = DiaCalendarioEscolar.Clasificacion.FERIADO
        martes.save()
        filas = previsualizar_asignacion(
            self.asignacion, self.calendario, date(2026, 8, 24), date(2026, 8, 28)
        )
        self.assertEqual(len([f for f in filas if f.estado == ProgramacionMenuEscolar.Estado.PROGRAMADO]), 4)
        self.assertEqual(filas[1].estado, ProgramacionMenuEscolar.Estado.SIN_DOCENCIA)
        self.assertEqual(filas[2].producto, "Muffin S1")

    def test_ciclo_semana_cinco_regresa_a_uno_por_semana_calendario(self):
        self.assertEqual(semana_ciclo(self.version, date(2026, 9, 21)), 5)
        self.assertEqual(semana_ciclo(self.version, date(2026, 9, 28)), 1)

    def test_semana_completa_sin_docencia_no_consume_ciclo(self):
        self.calendario.fin_docencia = date(2026, 9, 11)
        self.calendario.save(update_fields=("fin_docencia",))
        for desplazamiento in range(14, 19):
            DiaCalendarioEscolar.objects.create(
                calendario=self.calendario,
                fecha=self.calendario.inicio_docencia + timedelta(days=desplazamiento),
                clasificacion=DiaCalendarioEscolar.Clasificacion.DOCENCIA,
            )
        self.calendario.dias.filter(
            fecha__range=(date(2026, 8, 31), date(2026, 9, 4))
        ).update(clasificacion=DiaCalendarioEscolar.Clasificacion.NO_LECTIVO)

        self.assertEqual(
            semana_ciclo_escolar_regular(self.version, self.calendario, date(2026, 9, 7)),
            2,
        )

    def test_navidad_dos_semanas_completas_no_consume_ciclo(self):
        self.calendario.fin_docencia = date(2026, 10, 16)
        self.calendario.save(update_fields=("fin_docencia",))
        fecha = date(2026, 9, 7)
        while fecha <= self.calendario.fin_docencia:
            DiaCalendarioEscolar.objects.create(
                calendario=self.calendario,
                fecha=fecha,
                clasificacion=(
                    DiaCalendarioEscolar.Clasificacion.DOCENCIA
                    if fecha.weekday() < 5
                    else DiaCalendarioEscolar.Clasificacion.NO_LECTIVO
                ),
            )
            fecha += timedelta(days=1)
        self.calendario.dias.filter(
            fecha__range=(date(2026, 9, 21), date(2026, 10, 2))
        ).update(clasificacion=DiaCalendarioEscolar.Clasificacion.VACACIONES)

        self.assertEqual(semana_ciclo_escolar_regular(self.version, self.calendario, date(2026, 9, 14)), 4)
        self.assertEqual(semana_ciclo_escolar_regular(self.version, self.calendario, date(2026, 10, 5)), 5)
        self.assertEqual(semana_ciclo_escolar_regular(self.version, self.calendario, date(2026, 10, 12)), 1)

    def test_regreso_jueves_usa_jueves_y_viernes_de_semana_correspondiente(self):
        self.calendario.dias.filter(
            fecha__range=(date(2026, 8, 31), date(2026, 9, 2))
        ).update(clasificacion=DiaCalendarioEscolar.Clasificacion.NO_LECTIVO)
        filas = previsualizar_asignacion(
            self.asignacion, self.calendario, date(2026, 8, 31), date(2026, 9, 4)
        )

        self.assertEqual([fila.estado for fila in filas[:3]], ["SIN_DOCENCIA"] * 3)
        self.assertEqual(filas[3].semana, 2)
        self.assertEqual(filas[3].producto, "Pan integral S2")
        self.assertEqual(filas[4].producto, "Bizcocho S2")

    def test_version_vigente_no_se_usa_antes_de_fecha_operativa(self):
        self.version.fecha_uso_desde = date(2026, 8, 26)
        self.version.save(update_fields=("fecha_uso_desde",))

        self.assertIsNone(version_vigente(self.programa, date(2026, 8, 25)))
        self.assertEqual(version_vigente(self.programa, date(2026, 8, 26)), self.version)

    def test_version_historica_sin_fecha_uso_usa_vigencia_como_fallback(self):
        self.version.fecha_uso_desde = None
        self.version.save(update_fields=("fecha_uso_desde",))

        self.assertIsNone(version_vigente(self.programa, date(2026, 8, 23)))
        self.assertEqual(version_vigente(self.programa, date(2026, 8, 24)), self.version)

    def test_activacion_bloquea_version_activa_superpuesta(self):
        nueva = VersionProgramaMenu.objects.create(
            programa=self.programa,
            nombre="V2 superpuesta",
            vigente_desde=date(2026, 9, 1),
            fecha_uso_desde=date(2026, 9, 1),
            semanas_ciclo=5,
            fecha_ancla_ciclo=date(2026, 9, 1),
            estado=VersionProgramaMenu.Estado.BORRADOR,
        )
        for semana in range(1, 6):
            for dia in range(5):
                ItemCicloMenu.objects.create(
                    version=nueva, semana=semana, dia_semana=dia, producto="Pan"
                )
        self.client.force_login(self.user)

        response = self.client.post(reverse("activar_version_programa_menu", args=[nueva.pk]), follow=True)

        nueva.refresh_from_db()
        self.assertEqual(nueva.estado, VersionProgramaMenu.Estado.BORRADOR)
        self.assertContains(response, "se superpone con una versión activa")

    def _version_borrador_con_items(self, programa=None, nombre="Borrador eliminable"):
        version = VersionProgramaMenu.objects.create(
            programa=programa or self.programa,
            nombre=nombre,
            vigente_desde=date(2026, 10, 1),
            fecha_uso_desde=date(2026, 10, 1),
            semanas_ciclo=5,
            fecha_ancla_ciclo=date(2026, 10, 1),
            estado=VersionProgramaMenu.Estado.BORRADOR,
        )
        for semana in range(1, 6):
            for dia in range(5):
                ItemCicloMenu.objects.create(
                    version=version, semana=semana, dia_semana=dia, producto="Pan"
                )
        return version

    def test_elimina_version_borrador_y_sus_items(self):
        version = self._version_borrador_con_items()
        item_ids = list(version.items.values_list("pk", flat=True))
        self.client.force_login(self.user)

        response = self.client.post(reverse("eliminar_version_programa_menu", args=[version.pk]), follow=True)

        self.assertContains(response, "Versión eliminada correctamente")
        self.assertFalse(VersionProgramaMenu.objects.filter(pk=version.pk).exists())
        self.assertFalse(ItemCicloMenu.objects.filter(pk__in=item_ids).exists())

    def test_no_elimina_version_activa_y_requiere_post(self):
        self.client.force_login(self.user)
        url = reverse("eliminar_version_programa_menu", args=[self.version.pk])

        self.assertEqual(self.client.get(url).status_code, 405)
        response = self.client.post(url, follow=True)

        self.assertContains(response, "Solo se pueden eliminar versiones")
        self.assertTrue(VersionProgramaMenu.objects.filter(pk=self.version.pk).exists())

    def test_no_elimina_version_de_otra_empresa(self):
        otra_empresa = Empresa.objects.create(nombre="Empresa B")
        otro_programa = ProgramaMenu.objects.create(
            empresa=otra_empresa, codigo="MENU-B", nombre="Menú B",
            modalidad=ProgramaMenu.Modalidad.REGULAR,
        )
        version = self._version_borrador_con_items(otro_programa, "Borrador empresa B")
        self.client.force_login(self.user)

        response = self.client.post(reverse("eliminar_version_programa_menu", args=[version.pk]))

        self.assertEqual(response.status_code, 404)
        self.assertTrue(VersionProgramaMenu.objects.filter(pk=version.pk).exists())
        self.assertEqual(version.items.count(), 25)

    def test_historial_protegido_revierte_eliminacion_de_items(self):
        version = self._version_borrador_con_items(nombre="Borrador con historial")
        item = version.items.get(semana=1, dia_semana=0)
        ProgramacionMenuEscolar.objects.create(
            empresa=self.empresa,
            calendario=self.calendario,
            asignacion=self.asignacion,
            centro=self.centro,
            fecha=date(2026, 8, 24),
            dia_semana=0,
            semana_ciclo=1,
            programa=self.programa,
            version=version,
            item_ciclo=item,
            producto=item.producto,
            programa_snapshot=self.programa.nombre,
            version_snapshot=version.nombre,
            modalidad_snapshot=ProgramaMenu.Modalidad.REGULAR,
            estado=ProgramacionMenuEscolar.Estado.PROGRAMADO,
            creada_por=self.user,
        )
        self.client.force_login(self.user)

        response = self.client.post(reverse("eliminar_version_programa_menu", args=[version.pk]), follow=True)

        self.assertContains(response, "información histórica u operativa relacionada")
        self.assertTrue(VersionProgramaMenu.objects.filter(pk=version.pk).exists())
        self.assertEqual(ItemCicloMenu.objects.filter(version=version).count(), 25)

    def test_version_vigente_respeta_cambio_de_fecha(self):
        self.version.vigente_hasta = date(2026, 8, 31)
        self.version.save()
        v2 = VersionProgramaMenu.objects.create(
            programa=self.programa, nombre="V2", vigente_desde=date(2026, 9, 1),
            fecha_uso_desde=date(2026, 9, 1),
            semanas_ciclo=5, fecha_ancla_ciclo=date(2026, 9, 1), semana_inicial=1,
            estado=VersionProgramaMenu.Estado.ACTIVA,
        )
        for semana in range(1, 6):
            for dia in range(5):
                ItemCicloMenu.objects.create(version=v2, semana=semana, dia_semana=dia, producto="Pan V2")
        self.assertEqual(version_vigente(self.programa, date(2026, 8, 31)), self.version)
        self.assertEqual(version_vigente(self.programa, date(2026, 9, 1)), v2)

    def test_reiniciar_y_semana_especifica_son_deterministas(self):
        self.assertEqual(semana_ciclo(self.version, self.version.fecha_ancla_ciclo), 1)
        self.version.semana_inicial = 3
        self.version.modo_inicio_ciclo = VersionProgramaMenu.InicioCiclo.ESPECIFICA
        self.assertEqual(semana_ciclo(self.version, self.version.fecha_ancla_ciclo), 3)

    def test_materializacion_preserva_snapshot_historico(self):
        filas = materializar_programacion(
            self.asignacion, self.calendario, usuario=self.user,
            fecha_inicio=date(2026, 8, 24), fecha_fin=date(2026, 8, 24),
        )
        fila = filas[0]
        nombre = fila.version_snapshot
        self.version.nombre = "V1 renombrada"
        self.version.save()
        fila.refresh_from_db()
        self.assertEqual(fila.version_snapshot, nombre)

    def test_materializacion_no_genera_producto_en_fecha_excluida(self):
        dia = self.calendario.dias.get(fecha=date(2026, 8, 25))
        dia.clasificacion = DiaCalendarioEscolar.Clasificacion.SUSPENSION
        dia.save()
        filas = materializar_programacion(
            self.asignacion, self.calendario, usuario=self.user,
            fecha_inicio=dia.fecha, fecha_fin=dia.fecha,
        )
        self.assertEqual(filas[0].estado, ProgramacionMenuEscolar.Estado.SIN_DOCENCIA)
        self.assertEqual(filas[0].producto, "")

    def test_bebida_no_suministrada_no_genera_producto_operativo(self):
        item = self.version.items.get(semana=1, dia_semana=0)
        item.producto = "Leche"
        item.es_suministrado = False
        item.save()
        fila = previsualizar_asignacion(
            self.asignacion, self.calendario, date(2026, 8, 24), date(2026, 8, 24)
        )[0]
        self.assertEqual(fila.estado, ProgramacionMenuEscolar.Estado.NO_SUMINISTRADO)
        self.assertEqual(fila.producto, "")

    def _crear_prepara_visual(self, *, completo=True, codigo="PREPARA-VISUAL"):
        programa = ProgramaMenu.objects.create(
            empresa=self.empresa, codigo=codigo, nombre=f"Menú {codigo}",
            modalidad=ProgramaMenu.Modalidad.PREPARA,
        )
        version = VersionProgramaMenu.objects.create(
            programa=programa, nombre="2026-2028", vigente_desde=date(2026, 8, 1),
            fecha_uso_desde=date(2026, 8, 22), semanas_ciclo=5,
            fecha_ancla_ciclo=date(2026, 8, 22), estado=VersionProgramaMenu.Estado.ACTIVA,
        )
        posiciones = [(semana, dia) for semana in range(1, 6) for dia in (5, 6)]
        if not completo:
            posiciones.pop()
        for semana, dia in posiciones:
            ItemCicloMenu.objects.create(
                version=version, semana=semana, dia_semana=dia,
                producto="Pan" if semana % 2 else "Galleta de avena",
            )
        return version

    def test_prepara_completo_se_presenta_como_matriz_10_de_10(self):
        prepara = self._crear_prepara_visual()
        self.client.force_login(self.user)

        response = self.client.get(reverse("planificacion_menu_escolar"))
        prepara_contexto = next(
            version
            for programa in response.context["programas"]
            for version in programa.versiones.all()
            if version.pk == prepara.pk
        )

        self.assertEqual(prepara_contexto.total_posiciones_requeridas, 10)
        self.assertEqual(prepara_contexto.total_posiciones_configuradas, 10)
        self.assertTrue(prepara_contexto.ciclo_completo)
        self.assertEqual(len(prepara_contexto.matriz_ciclo), 5)
        self.assertContains(response, "Sábado")
        self.assertContains(response, "Domingo")
        self.assertNotContains(response, "Semana 1 · Sábado:")

    def test_prepara_incompleto_no_es_ciclo_completo_y_regular_no_cambia(self):
        prepara = self._crear_prepara_visual(completo=False, codigo="PREPARA-INCOMPLETO")
        self.client.force_login(self.user)

        response = self.client.get(reverse("planificacion_menu_escolar"))
        versiones = {
            version.pk: version
            for programa in response.context["programas"]
            for version in programa.versiones.all()
        }

        self.assertEqual(versiones[prepara.pk].total_posiciones_requeridas, 10)
        self.assertEqual(versiones[prepara.pk].total_posiciones_configuradas, 9)
        self.assertFalse(versiones[prepara.pk].ciclo_completo)
        self.assertEqual(versiones[self.version.pk].total_posiciones_requeridas, 25)
        self.assertEqual(versiones[self.version.pk].total_posiciones_configuradas, 25)
        self.assertTrue(versiones[self.version.pk].ciclo_completo)
        for encabezado in ("Lunes", "Martes", "Miércoles", "Jueves", "Viernes"):
            self.assertContains(response, encabezado)

    def test_prepara_es_independiente_y_no_desplaza_sabado_a_domingo(self):
        prepara = ProgramaMenu.objects.create(
            empresa=self.empresa, codigo="PREPARA", nombre="PAE PREPARA", modalidad=ProgramaMenu.Modalidad.PREPARA
        )
        version = VersionProgramaMenu.objects.create(
            programa=prepara, nombre="P1", vigente_desde=date(2026, 8, 24), semanas_ciclo=2,
            fecha_ancla_ciclo=date(2026, 8, 24), estado=VersionProgramaMenu.Estado.ACTIVA,
        )
        ItemCicloMenu.objects.create(version=version, semana=1, dia_semana=5, producto="Pan sabado")
        ItemCicloMenu.objects.create(version=version, semana=1, dia_semana=6, producto="Galleta domingo")
        asignacion = AsignacionProgramaCentro.objects.create(
            centro=self.centro, programa=prepara, modalidad=ProgramaMenu.Modalidad.PREPARA,
            dias_entrega=[5, 6], vigente_desde=date(2026, 8, 24),
        )
        sabado = self.calendario.dias.get(fecha=date(2026, 8, 29))
        domingo = self.calendario.dias.get(fecha=date(2026, 8, 30))
        sabado.clasificacion = DiaCalendarioEscolar.Clasificacion.FERIADO; sabado.save()
        domingo.clasificacion = DiaCalendarioEscolar.Clasificacion.DOCENCIA; domingo.save()
        filas = previsualizar_asignacion(asignacion, self.calendario, sabado.fecha, domingo.fecha)
        self.assertEqual(filas[0].estado, ProgramacionMenuEscolar.Estado.SIN_DOCENCIA)
        self.assertEqual(filas[1].producto, "Galleta domingo")
        self.assertEqual(contar_docencia_regular(self.calendario), 10)

    def test_asignacion_rechaza_programa_de_otra_empresa(self):
        otro = Empresa.objects.create(nombre="Empresa B")
        programa = ProgramaMenu.objects.create(empresa=otro, codigo="B", nombre="B", modalidad="REGULAR")
        asignacion = AsignacionProgramaCentro(
            centro=self.centro, programa=programa, modalidad="REGULAR", dias_entrega=[0],
            vigente_desde=date(2026, 8, 24),
        )
        with self.assertRaises(ValidationError):
            asignacion.full_clean()

    def test_selector_solo_ofrece_programas_con_version_activa(self):
        borrador = ProgramaMenu.objects.create(
            empresa=self.empresa, codigo="BORRADOR", nombre="Solo borrador",
            modalidad=ProgramaMenu.Modalidad.REGULAR,
        )
        VersionProgramaMenu.objects.create(
            programa=borrador, nombre="B1", vigente_desde=date(2026, 8, 24),
            semanas_ciclo=5, fecha_ancla_ciclo=date(2026, 8, 24),
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse("planificacion_menu_escolar"))
        asignables = list(response.context["programas_asignables"])

        self.assertIn(self.programa, asignables)
        self.assertNotIn(borrador, asignables)

    def test_backend_rechaza_programa_sin_version_activa(self):
        borrador = ProgramaMenu.objects.create(
            empresa=self.empresa, codigo="BORRADOR", nombre="Solo borrador",
            modalidad=ProgramaMenu.Modalidad.REGULAR,
        )
        VersionProgramaMenu.objects.create(
            programa=borrador, nombre="B1", vigente_desde=date(2026, 8, 24),
            semanas_ciclo=5, fecha_ancla_ciclo=date(2026, 8, 24),
        )
        self.client.force_login(self.user)

        response = self.client.post(reverse("asignar_programa_centro"), {
            "centro_id": self.centro.pk, "programa_id": borrador.pk,
            "dias_entrega": ["0"], "vigente_desde": "2026-09-01",
        }, follow=True)

        self.assertContains(response, "no tiene ninguna versión activa")
        self.assertFalse(AsignacionProgramaCentro.objects.filter(programa=borrador).exists())

    def test_editar_asignacion_propia_no_reescribe_historial(self):
        historica = materializar_programacion(
            self.asignacion, self.calendario, usuario=self.user,
            fecha_inicio=date(2026, 8, 24), fecha_fin=date(2026, 8, 24),
        )[0]
        snapshot = (historica.centro_id, historica.programa_id, historica.producto, historica.version_snapshot)
        self.client.force_login(self.user)

        response = self.client.post(reverse("editar_asignacion_programa_centro", args=[self.asignacion.pk]), {
            "centro_id": self.centro.pk, "programa_id": self.programa.pk,
            "dias_entrega": ["1", "3"], "vigente_desde": "2026-09-01",
            "vigente_hasta": "2027-06-18",
        })

        self.assertRedirects(response, reverse("planificacion_menu_escolar"))
        self.asignacion.refresh_from_db(); historica.refresh_from_db()
        self.assertEqual(self.asignacion.dias_entrega, [1, 3])
        self.assertEqual(self.asignacion.vigente_desde, date(2026, 9, 1))
        self.assertEqual((historica.centro_id, historica.programa_id, historica.producto, historica.version_snapshot), snapshot)

    def test_editar_rechaza_centro_programa_ajenos_y_programa_borrador(self):
        otro = Empresa.objects.create(nombre="Empresa ajena")
        centro_ajeno = CentroEducativo.objects.create(empresa=otro, codigo="B-1", nombre="Centro B")
        programa_ajeno = ProgramaMenu.objects.create(
            empresa=otro, codigo="PB", nombre="Programa B", modalidad=ProgramaMenu.Modalidad.REGULAR,
        )
        borrador = ProgramaMenu.objects.create(
            empresa=self.empresa, codigo="DRAFT", nombre="Draft", modalidad=ProgramaMenu.Modalidad.REGULAR,
        )
        self.client.force_login(self.user)
        url = reverse("editar_asignacion_programa_centro", args=[self.asignacion.pk])
        base = {"dias_entrega": ["0"], "vigente_desde": "2026-09-01"}

        self.assertEqual(self.client.post(url, {**base, "centro_id": centro_ajeno.pk, "programa_id": self.programa.pk}).status_code, 404)
        self.assertEqual(self.client.post(url, {**base, "centro_id": self.centro.pk, "programa_id": programa_ajeno.pk}).status_code, 404)
        response = self.client.post(url, {**base, "centro_id": self.centro.pk, "programa_id": borrador.pk}, follow=True)
        self.assertContains(response, "no tiene ninguna versión activa")
        self.asignacion.refresh_from_db()
        self.assertEqual(self.asignacion.programa, self.programa)

    def test_no_edita_ni_elimina_asignacion_de_otra_empresa(self):
        otro_usuario = get_user_model().objects.create_user("otro-admin", password="x")
        otra_empresa = Empresa.objects.create(usuario=otro_usuario, nombre="Empresa B", modulo_menu=True)
        otro_centro = CentroEducativo.objects.create(empresa=otra_empresa, codigo="B-1", nombre="Centro B")
        otro_programa = ProgramaMenu.objects.create(
            empresa=otra_empresa, codigo="PB", nombre="Programa B", modalidad=ProgramaMenu.Modalidad.REGULAR,
        )
        otra_asignacion = AsignacionProgramaCentro.objects.create(
            centro=otro_centro, programa=otro_programa, modalidad=ProgramaMenu.Modalidad.REGULAR,
            dias_entrega=[0], vigente_desde=date(2026, 8, 24),
        )
        self.client.force_login(self.user)

        self.assertEqual(self.client.get(reverse("editar_asignacion_programa_centro", args=[otra_asignacion.pk])).status_code, 404)
        self.assertEqual(self.client.post(reverse("eliminar_asignacion_programa_centro", args=[otra_asignacion.pk])).status_code, 404)
        self.assertTrue(AsignacionProgramaCentro.objects.filter(pk=otra_asignacion.pk).exists())

    def test_eliminar_asignacion_propia_requiere_post(self):
        url = reverse("eliminar_asignacion_programa_centro", args=[self.asignacion.pk])
        self.client.force_login(self.user)

        self.assertEqual(self.client.get(url).status_code, 405)
        response = self.client.post(url)

        self.assertRedirects(response, reverse("planificacion_menu_escolar"))
        self.assertFalse(AsignacionProgramaCentro.objects.filter(pk=self.asignacion.pk).exists())

    def test_eliminar_asignacion_con_historial_es_protegido(self):
        materializar_programacion(
            self.asignacion, self.calendario, usuario=self.user,
            fecha_inicio=date(2026, 8, 24), fecha_fin=date(2026, 8, 24),
        )
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("eliminar_asignacion_programa_centro", args=[self.asignacion.pk]), follow=True
        )

        self.assertContains(response, "posee programación histórica")
        self.assertTrue(AsignacionProgramaCentro.objects.filter(pk=self.asignacion.pk).exists())
        self.assertEqual(ProgramacionMenuEscolar.objects.filter(asignacion=self.asignacion).count(), 1)

    def test_preview_rechaza_calendario_de_otra_empresa(self):
        otro = Empresa.objects.create(nombre="Empresa B")
        calendario = CalendarioEscolar.objects.create(
            empresa=otro, nombre="B", anio_inicio=2026, anio_fin=2027,
            inicio_docencia=date(2026, 8, 24), fin_docencia=date(2026, 8, 25),
        )
        with self.assertRaises(ValidationError):
            previsualizar_asignacion(self.asignacion, calendario)

    def test_override_auditable_y_bloqueo_historico(self):
        fila = materializar_programacion(
            self.asignacion, self.calendario, usuario=self.user,
            fecha_inicio=date(2026, 8, 24), fecha_fin=date(2026, 8, 24),
        )[0]
        excepcion = aplicar_excepcion(
            fila, usuario=self.user, estado_nuevo=ProgramacionMenuEscolar.Estado.EXTRAORDINARIO,
            producto_nuevo="Pan especial", motivo="Orden extraordinaria",
        )
        self.assertEqual(excepcion.producto_anterior, "Pan S1")
        fila.refresh_from_db(); self.assertEqual(fila.producto, "Pan especial")
        fila.bloqueada = True; fila.save()
        with self.assertRaises(ValidationError):
            aplicar_excepcion(fila, usuario=self.user, estado_nuevo="OMITIDO", motivo="Prueba")

    def test_exportaciones_excel_pdf_y_rango(self):
        materializar_programacion(
            self.asignacion, self.calendario, usuario=self.user,
            fecha_inicio=date(2026, 8, 24), fecha_fin=date(2026, 8, 25),
        )
        self.client.force_login(self.user)
        excel = self.client.get(reverse("exportar_programacion_excel"), {"fecha_inicio": "2026-08-25", "fecha_fin": "2026-08-25"})
        self.assertEqual(excel.status_code, 200)
        ws = load_workbook(BytesIO(excel.content)).active

        # El Excel ejecutivo reserva las primeras filas para identidad
        # empresarial, título, período y datos de generación.
        self.assertEqual(ws.title, "Programación menú")
        self.assertEqual(ws["A5"].value, "PROGRAMACIÓN DE MENÚ ESCOLAR")
        self.assertEqual(ws["A9"].value, "Fecha")
        self.assertEqual(ws["B9"].value, "Día")
        self.assertEqual(ws["D9"].value, "Centro educativo")
        self.assertEqual(ws["J9"].value, "Estado")

        # El filtro solicitado debe devolver únicamente el 25/08/2026.
        self.assertEqual(ws.max_row, 10)
        self.assertEqual(ws["A10"].value.strftime("%d/%m/%Y"), "25/08/2026")
        self.assertEqual(ws["B10"].value, "Martes")

        pdf = self.client.get(reverse("exportar_programacion_pdf"), {"fecha_inicio": "2026-08-25", "fecha_fin": "2026-08-25"})
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(b"".join(pdf.streaming_content)[:4], b"%PDF")

    def test_usuario_no_autorizado_no_modifica_y_empresa_ajena_no_se_exporta(self):
        otro_usuario = get_user_model().objects.create_user("consulta")
        otra_empresa = Empresa.objects.create(usuario=otro_usuario, nombre="Empresa B", modulo_menu=True)
        self.client.force_login(otro_usuario)
        response = self.client.post(reverse("crear_programa_menu"), {"codigo": "X", "nombre": "X", "modalidad": "REGULAR"})
        self.assertEqual(response.status_code, 403)
        response = self.client.get(reverse("exportar_programacion_excel"))
        self.assertEqual(response.status_code, 200)
        self.assertNotIn(b"Panaderia A", response.content)
        self.assertNotEqual(otra_empresa.pk, self.empresa.pk)
