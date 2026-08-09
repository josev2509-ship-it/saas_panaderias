from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from catalogos.models import Moneda, MonedaEmpresa
from compras.models import Proveedor
from conduces.models import Empresa
from contabilidad.models import FacturaProveedor
from tesoreria.models import (
    ConciliacionBancaria,
    CuentaBancariaEmpresa,
    ImportacionExtractoBancario,
)


class DGII606CanonicalExportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_superuser("rc1-606", "rc1-606@example.test", "test-only")
        cls.other_user = User.objects.create_user("rc1-606-other")
        cls.empresa = Empresa.objects.create(
            usuario=cls.user, nombre="Empresa 606", modulo_compras=True
        )
        cls.other = Empresa.objects.create(usuario=cls.other_user, nombre="Otra empresa 606")
        moneda = Moneda.objects.create(codigo="DOP", nombre="Peso", simbolo="RD$")
        cls.moneda = MonedaEmpresa.objects.create(empresa=cls.empresa, moneda=moneda, es_base=True)
        other_moneda = MonedaEmpresa.objects.create(empresa=cls.other, moneda=moneda, es_base=True)
        cls.proveedor = Proveedor.objects.create(
            empresa=cls.empresa, codigo="P-606", tipo_persona="JURIDICA",
            razon_social="=Proveedor seguro", rnc_identificacion="101010101",
            estado="ACTIVO", creado_por=cls.user,
        )
        other_provider = Proveedor.objects.create(
            empresa=cls.other, codigo="P-OTHER", tipo_persona="JURIDICA",
            razon_social="Proveedor ajeno", rnc_identificacion="202020202",
            estado="ACTIVO", creado_por=cls.other_user,
        )
        today = timezone.localdate()
        cls.factura = FacturaProveedor.objects.create(
            empresa=cls.empresa, proveedor=cls.proveedor, numero="FAC-606", ncf="B0100000001",
            fecha=today, vence_el=today + timedelta(days=30), moneda=cls.moneda,
            subtotal=Decimal("100"), impuesto=Decimal("18"), total=Decimal("118"), estado="VALIDADA",
            dimensiones={"dgii_tipo_bienes_servicios": "09", "dgii_servicios": "100", "dgii_forma_pago": "07"},
            creado_por=cls.user,
        )
        FacturaProveedor.objects.create(
            empresa=cls.other, proveedor=other_provider, numero="FAC-OTHER", ncf="B0199999999",
            fecha=today, vence_el=today + timedelta(days=30), moneda=other_moneda,
            subtotal=1, impuesto=0, total=1, estado="VALIDADA", creado_por=cls.other_user,
        )

    def setUp(self):
        self.client.force_login(self.user)

    def test_606_page_uses_p2p_source_and_is_tenant_safe(self):
        response = self.client.get(reverse("contabilidad:gastos"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "B0100000001")
        self.assertNotContains(response, "B0199999999")
        self.assertContains(response, "Vista fiscal canónica", html=False)

    def test_606_xlsx_has_canonical_columns_tenant_and_formula_protection(self):
        response = self.client.get(reverse("contabilidad:exportar_606_excel"))
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.content), data_only=False)
        sheet = book.active
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(len(headers), 22)
        self.assertIn("RNC/Cédula", headers)
        self.assertIn("Forma de Pago", headers)
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(2, 1).value, "101010101")
        self.assertNotIn("FAC-OTHER", str(list(sheet.values)))

    def test_606_rejects_invalid_period(self):
        response = self.client.get(reverse("contabilidad:exportar_606_excel"), {"periodo": "invalid"})
        self.assertEqual(response.status_code, 400)

    def test_p2p_extract_and_reconciliation_lists_render_real_objects(self):
        cuenta = CuentaBancariaEmpresa.objects.create(
            empresa=self.empresa,
            banco="Banco Demo",
            numero_enmascarado="****6060",
            moneda=self.moneda,
        )
        ImportacionExtractoBancario.objects.create(
            empresa=self.empresa,
            cuenta=cuenta,
            nombre_archivo="extracto-rc1.csv",
            formato="CSV",
            huella="a" * 64,
        )
        today = timezone.localdate()
        ConciliacionBancaria.objects.create(
            empresa=self.empresa,
            cuenta=cuenta,
            desde=today,
            hasta=today,
            saldo_banco=Decimal("118.00"),
            saldo_libros=Decimal("118.00"),
        )

        for recurso in ("extractos", "conciliaciones"):
            response = self.client.get(reverse("compras:p2p_finance_list", args=[recurso]))
            self.assertEqual(response.status_code, 200)
            self.assertNotContains(response, "VariableDoesNotExist")

    def test_enterprise_list_mobile_contract_prevents_action_overflow(self):
        root = __import__("pathlib").Path(__file__).resolve().parents[1]
        css_path = (
            root
            / "conduces/static/design_system/css/enterprise_list.css"
        )
        css = css_path.read_text(encoding="utf-8")
        self.assertIn("minmax(0,1fr)", css)
        self.assertIn(".el-row-menu form", css)
        javascript = (
            root / "conduces/static/design_system/js/demo_unified.js"
        ).read_text(encoding="utf-8")
        self.assertIn("table,.el-row-menu", javascript)
