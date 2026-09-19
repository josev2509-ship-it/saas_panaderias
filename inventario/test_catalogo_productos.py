from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from conduces.models import Empresa
from inventario.models import MovimientoInventario, ProductoInventario


class CatalogoInventarioTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("catalogo", password="x")
        self.empresa = Empresa.objects.create(usuario=self.user, nombre="Empresa Catalogo", activa=True)
        self.client.force_login(self.user)

    def test_crear_producto_manual_registra_stock_inicial_como_movimiento(self):
        respuesta = self.client.post(reverse("inventario:crear_producto_inventario"), {
            "codigo": "HAR-F", "nombre": "Harina fuerte", "tipo": "materia_prima",
            "clasificacion_operativa": "materia_prima", "unidad_medida": "lb",
            "unidad_compra": "saco", "cantidad_por_empaque": "120", "stock_inicial": "240",
            "stock_minimo": "120", "precio_unitario_compra": "2450", "porcentaje_itbis": "0",
            "proveedor": "COOPROHARINA", "afecta_produccion": "on", "activo": "on",
        })
        self.assertEqual(respuesta.status_code, 302)
        producto = ProductoInventario.objects.get(empresa=self.empresa, codigo="HAR-F")
        self.assertEqual(producto.stock_actual, Decimal("240.0000"))
        self.assertEqual(producto.cantidad_por_empaque, Decimal("120.0000"))
        self.assertTrue(producto.afecta_produccion)
        movimiento = MovimientoInventario.objects.get(empresa=self.empresa, producto=producto)
        self.assertEqual(movimiento.tipo, "ajuste")
        self.assertEqual(movimiento.saldo_posterior, Decimal("240.0000"))

    def test_editar_producto_guarda_clasificacion_y_afecta_produccion(self):
        producto = ProductoInventario.objects.create(empresa=self.empresa, codigo="EMP-1", nombre="Funda", tipo="empaque", unidad_medida="unidad", clasificacion_operativa="empaque", afecta_produccion=False)
        respuesta = self.client.post(reverse("inventario:editar_producto_inventario", args=[producto.pk]), {
            "codigo": "EMP-1", "nombre": "Funda 14 x 25", "tipo": "empaque",
            "clasificacion_operativa": "empaque", "unidad_medida": "unidad", "unidad_compra": "fardo",
            "cantidad_por_empaque": "100", "stock_minimo": "0", "precio_unitario_compra": "1634.94",
            "porcentaje_itbis": "18", "proveedor": "COOPROHARINA", "activo": "on",
        })
        self.assertEqual(respuesta.status_code, 302)
        producto.refresh_from_db()
        self.assertEqual(producto.clasificacion_operativa, "empaque")
        self.assertFalse(producto.afecta_produccion)
        self.assertEqual(producto.unidad_compra, "fardo")

    def test_descargar_excel_incluye_catalogo_actual(self):
        ProductoInventario.objects.create(empresa=self.empresa, codigo="HAR-N", nombre="Harina normal", tipo="materia_prima", unidad_medida="lb", unidad_compra="saco", cantidad_por_empaque=100, precio_unitario_compra=2177.39, clasificacion_operativa="materia_prima", afecta_produccion=True)
        respuesta = self.client.get(reverse("inventario:descargar_plantilla"))
        self.assertEqual(respuesta.status_code, 200)
        libro = load_workbook(BytesIO(respuesta.content), data_only=True)
        hoja = libro.active
        encabezados = [c.value for c in hoja[1]]
        self.assertIn("clasificacion_operativa", encabezados)
        self.assertIn("afecta_produccion", encabezados)
        valores = list(hoja.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(valores), 1)
        self.assertIn("Harina normal", valores[0])

    def test_plantilla_vacia_no_exporta_productos(self):
        ProductoInventario.objects.create(empresa=self.empresa, codigo="HAR-X", nombre="Harina X", tipo="materia_prima")
        respuesta = self.client.get(reverse("inventario:descargar_plantilla") + "?vacia=1")
        libro = load_workbook(BytesIO(respuesta.content), data_only=True)
        self.assertEqual(libro.active.max_row, 1)

    def test_excel_actualiza_producto_provisional_por_nombre_sin_duplicar(self):
        ProductoInventario.objects.create(empresa=self.empresa, codigo=None, nombre="Harina fuerte", tipo="materia_prima", unidad_medida="lb")
        libro = Workbook(); hoja = libro.active
        hoja.append(["codigo", "nombre", "tipo", "clasificacion_operativa", "afecta_produccion", "unidad_medida", "unidad_compra", "cantidad_por_empaque", "stock_actual", "stock_minimo", "precio_unitario_compra", "porcentaje_itbis", "proveedor", "activo"])
        hoja.append(["PI-0004", "Harina fuerte", "materia_prima", "materia_prima", True, "lb", "saco", 120, 0, 0, 2450, 0, "COOPROHARINA", True])
        salida = BytesIO(); libro.save(salida); salida.seek(0); salida.name = "inventario.xlsx"
        respuesta = self.client.post(reverse("inventario:cargar_excel"), {"archivo": salida})
        self.assertEqual(respuesta.status_code, 302)
        self.assertEqual(ProductoInventario.objects.filter(empresa=self.empresa, nombre__iexact="Harina fuerte").count(), 1)
        producto = ProductoInventario.objects.get(empresa=self.empresa, nombre__iexact="Harina fuerte")
        self.assertEqual(producto.codigo, "PI-0004")
        self.assertEqual(producto.cantidad_por_empaque, Decimal("120.0000"))
