from io import BytesIO
from django.utils import timezone

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import Table, TableStyle, Paragraph
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from django.db.models import Sum
from django.db.models import Sum
from decimal import Decimal

from decimal import Decimal, InvalidOperation, ROUND_CEILING
from django.shortcuts import render, redirect, get_object_or_404

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.db import models
from django.http import HttpResponse
from django.contrib import messages
from core.application.operation_context import OperationContext
from core.application.request_idempotency import resolve_idempotency_context
from .engine import InventoryEngine

from openpyxl import Workbook, load_workbook

from conduces.models import MenuDiario, CentroEducativo

from django.db.models import Sum
from reportlab.platypus import Table, TableStyle
from reportlab.lib import colors
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter

from .models import (
    ProyeccionConsumoDiario,
    DetalleConsumoProduccion,
)


def _movimiento_desde_vista(request, *, empresa, producto, tipo, cantidad,
                            referencia="", observacion="", **kwargs):
    operation = f"inventario.views.{tipo}"
    context = OperationContext(
        empresa=empresa, usuario=request.user, request=request,
        observaciones=observacion,
        **resolve_idempotency_context(
            request, operation=operation, reference=referencia,
            metadata={"product_id": producto.pk},
        ),
    )
    return InventoryEngine.apply_movement(
        context=context, producto=producto, tipo=tipo, cantidad=cantidad,
        referencia=referencia,
    )


def _bloquear_mutacion_orden_legada(view_func):
    """Mantiene URLs históricas sin permitir nuevas mutaciones del flujo legado."""
    def wrapper(request, *args, **kwargs):
        if not request.user.is_authenticated:
            from django.contrib.auth.views import redirect_to_login
            return redirect_to_login(request.get_full_path())
        return render(request, "inventario/orden_compra_legada_bloqueada.html", status=403)
    wrapper.__name__ = view_func.__name__
    wrapper.__doc__ = view_func.__doc__
    return wrapper

from .models import (
    ProductoInventario,
    ProductoProduccion,
    Receta,
    DetalleReceta,
    MovimientoInventario,
    PrestamoMateriaPrima,
    OrdenCompra,
    DetalleOrdenCompra,
    ProduccionProgramada,
)


def obtener_empresa_usuario(request):
    if hasattr(request.user, "empresa_principal"):
        return request.user.empresa_principal
    return None


def convertir_decimal(valor, defecto="0"):
    if valor is None or valor == "":
        return Decimal(defecto)

    try:
        valor_limpio = str(valor).strip()
        valor_limpio = valor_limpio.replace("RD$", "")
        valor_limpio = valor_limpio.replace("$", "")
        valor_limpio = valor_limpio.replace("%", "")
        valor_limpio = valor_limpio.replace(",", "")
        return Decimal(valor_limpio)
    except (InvalidOperation, ValueError):
        return Decimal(defecto)


def convertir_itbis(valor):
    itbis = convertir_decimal(valor, "0")

    if itbis > 0 and itbis < 1:
        itbis = itbis * 100

    return itbis


def normalizar_texto(texto):
    return str(texto or "").strip().lower()


def buscar_receta_por_producto(empresa, producto_menu):
    producto_menu = normalizar_texto(producto_menu)

    if not producto_menu:
        return None

    receta = Receta.objects.filter(
        empresa=empresa,
        activa=True,
        producto__nombre__icontains=producto_menu
    ).first()

    if receta:
        return receta

    receta = Receta.objects.filter(
        empresa=empresa,
        activa=True,
        nombre__icontains=producto_menu
    ).first()

    if receta:
        return receta

    if "muffin" in producto_menu or "mufin" in producto_menu or "bizcocho" in producto_menu:
        return Receta.objects.filter(
            empresa=empresa,
            activa=True,
            producto__tipo="bizcocho"
        ).first()

    if "galleta" in producto_menu:
        return Receta.objects.filter(
            empresa=empresa,
            activa=True,
            producto__tipo="galleta"
        ).first()

    if "zanahoria" in producto_menu and "pan" in producto_menu:
        return Receta.objects.filter(
            empresa=empresa,
            activa=True,
            producto__tipo="pan_vegetales"
        ).first()

    if "pan" in producto_menu:
        return Receta.objects.filter(
            empresa=empresa,
            activa=True,
            producto__tipo="pan"
        ).first()

    return None


def cantidad_empaques_sugeridos(cantidad_faltante, cantidad_por_empaque):
    cantidad_faltante = Decimal(cantidad_faltante or 0)
    cantidad_por_empaque = Decimal(cantidad_por_empaque or 1)

    if cantidad_faltante <= 0:
        return Decimal("0")

    if cantidad_por_empaque <= 0:
        cantidad_por_empaque = Decimal("1")

    return (cantidad_faltante / cantidad_por_empaque).to_integral_value(
        rounding=ROUND_CEILING
    )


# =====================================================
# DASHBOARD
# =====================================================

# =====================================================
# DASHBOARD
# =====================================================

def generar_proyeccion_automatica_empresa(empresa, usuario=None):
    hoy = timezone.localdate()

    menus = MenuDiario.objects.filter(
        empresa=empresa,
        fecha__gte=hoy
    ).order_by("fecha")

    if not menus.exists():
        return

    fecha_fin = menus.last().fecha

    centros = CentroEducativo.objects.filter(
        empresa=empresa
    )

    total_raciones = sum(
        Decimal(centro.matricula or 0)
        for centro in centros
    )

    if total_raciones <= 0:
        return

    ProyeccionConsumoDiario.objects.filter(
        empresa=empresa,
        fecha__range=[hoy, fecha_fin]
    ).delete()

    for menu in menus:
        receta = buscar_receta_por_producto(
            empresa,
            menu.producto
        )

        if not receta:
            continue

        if not receta.rendimiento_unidades:
            continue

        for detalle in receta.detalles.all():
            cantidad_por_racion = (
                Decimal(detalle.cantidad or 0)
                / Decimal(receta.rendimiento_unidades or 1)
            )

            cantidad_total = cantidad_por_racion * total_raciones

            if receta.porcentaje_variacion:
                cantidad_total = cantidad_total * (
                    Decimal("1")
                    + (Decimal(receta.porcentaje_variacion or 0) / Decimal("100"))
                )

            ProyeccionConsumoDiario.objects.update_or_create(
                empresa=empresa,
                fecha=menu.fecha,
                producto_menu=menu.producto,
                materia_prima=detalle.materia_prima,
                defaults={
                    "receta": receta,
                    "raciones": total_raciones,
                    "cantidad_proyectada": cantidad_total,
                    "generado_por": usuario,
                }
            )


@login_required
def dashboard_inventario(request):
    empresa = obtener_empresa_usuario(request)

    if empresa:
        generar_proyeccion_automatica_empresa(
            empresa=empresa,
            usuario=request.user
        )

    productos = (
        ProductoInventario.objects.filter(
            empresa=empresa,
            activo=True
        )
        if empresa
        else ProductoInventario.objects.none()
    )

    total_productos = productos.count()

    productos_bajo_minimo_qs = productos.filter(
        stock_actual__lte=models.F("stock_minimo")
    )

    valor_total = sum(
        producto.valor_actual()
        for producto in productos
    )

    hoy = timezone.localdate()

    proyecciones = (
        ProyeccionConsumoDiario.objects
        .filter(
            empresa=empresa,
            fecha__gte=hoy
        )
        .order_by("fecha")
        if empresa
        else ProyeccionConsumoDiario.objects.none()
    )

    stock_virtual = {
        producto.id: Decimal(producto.stock_actual or 0)
        for producto in productos
    }

    fecha_disponible_hasta = None
    fecha_alerta_compra = None
    dias_estimados_produccion = 0
    producto_limitante = None

    fechas = proyecciones.values_list(
        "fecha",
        flat=True
    ).distinct()

    for fecha in fechas:
        consumos_dia = proyecciones.filter(
    fecha=fecha,
    materia_prima__afecta_produccion=True
)

        puede_producir = True

        for consumo in consumos_dia:
            producto_id = consumo.materia_prima_id
            cantidad = Decimal(consumo.cantidad_proyectada or 0)

            disponible = stock_virtual.get(
                producto_id,
                Decimal("0")
            )

            if disponible < cantidad:
                puede_producir = False
                producto_limitante = consumo.materia_prima
                break

        if not puede_producir:
            break

        for consumo in consumos_dia:
            producto_id = consumo.materia_prima_id
            cantidad = Decimal(consumo.cantidad_proyectada or 0)

            stock_virtual[producto_id] = (
                stock_virtual.get(producto_id, Decimal("0"))
                - cantidad
            )

        fecha_disponible_hasta = fecha
        dias_estimados_produccion += 1

    if fecha_disponible_hasta:
        fecha_alerta_compra = (
            fecha_disponible_hasta
            - timezone.timedelta(days=3)
        )

    materias_criticas = []

    for producto in productos_bajo_minimo_qs[:20]:
        materias_criticas.append({
            "producto": producto,
            "dias_restantes": 0,
        })

    contexto = {
        "titulo": "Dashboard Inventario",
        "total_productos": total_productos,
        "productos_bajo_minimo": productos_bajo_minimo_qs.count(),
        "valor_total": valor_total,
        "materias_criticas": materias_criticas,
        "fecha_disponible_hasta": fecha_disponible_hasta,
        "fecha_alerta_compra": fecha_alerta_compra,
        "dias_estimados_produccion": dias_estimados_produccion,
        "producto_limitante": producto_limitante,
    }

    return render(
        request,
        "inventario/dashboard.html",
        contexto
    )
# =====================================================
# PRÉSTAMOS DE MATERIA PRIMA
# =====================================================

@login_required
def crear_prestamo(request):
    empresa = obtener_empresa_usuario(request)

    productos = ProductoInventario.objects.filter(
        empresa=empresa,
        activo=True
    ).order_by("nombre")

    if request.method == "POST":
        producto = get_object_or_404(
            ProductoInventario,
            id=request.POST.get("producto_id"),
            empresa=empresa
        )

        tipo = request.POST.get("tipo")
        cantidad = convertir_decimal(request.POST.get("cantidad"), "0")

        prestamo = PrestamoMateriaPrima.objects.create(
            empresa=empresa,
            producto=producto,
            tipo=tipo,
            tercero=request.POST.get("tercero", "").strip(),
            cantidad=cantidad,
            fecha_prestamo=request.POST.get("fecha_prestamo") or timezone.localdate(),
            fecha_compromiso=request.POST.get("fecha_compromiso") or None,
            responsable_entrega=request.POST.get("responsable_entrega", "").strip(),
            responsable_recibe=request.POST.get("responsable_recibe", "").strip(),
            observacion=request.POST.get("observacion", "").strip(),
            usuario=request.user,
        )

        if tipo == "entregado":
            _movimiento_desde_vista(request,
                empresa=empresa,
                producto=producto,
                tipo="prestamo_entregado",
                cantidad=cantidad,
                referencia=f"PREST-{prestamo.id}",
                observacion=f"Préstamo entregado a {prestamo.tercero}",
                usuario=request.user,
            )

        if tipo == "recibido":
            _movimiento_desde_vista(request,
                empresa=empresa,
                producto=producto,
                tipo="prestamo_recibido",
                cantidad=cantidad,
                referencia=f"PREST-{prestamo.id}",
                observacion=f"Préstamo recibido de {prestamo.tercero}",
                usuario=request.user,
            )

        messages.success(request, "Préstamo registrado correctamente.")
        return redirect("inventario:prestamos")

    return render(request, "inventario/crear_prestamo.html", {
        "productos": productos,
    })


@login_required
def registrar_devolucion_prestamo(request, prestamo_id):
    empresa = obtener_empresa_usuario(request)

    prestamo = get_object_or_404(
        PrestamoMateriaPrima,
        id=prestamo_id,
        empresa=empresa
    )

    if request.method == "POST":
        cantidad = convertir_decimal(request.POST.get("cantidad_devuelta"), "0")

        if cantidad <= 0:
            messages.error(request, "La cantidad devuelta debe ser mayor que cero.")
            return redirect("inventario:prestamos")

        prestamo.cantidad_devuelta += cantidad
        prestamo.actualizar_estado()

        if prestamo.tipo == "entregado":
            tipo_movimiento = "devolucion_prestamo"
        else:
            tipo_movimiento = "salida"

        _movimiento_desde_vista(request,
            empresa=empresa,
            producto=prestamo.producto,
            tipo=tipo_movimiento,
            cantidad=cantidad,
            referencia=f"DEV-PREST-{prestamo.id}",
            observacion=f"Devolución de préstamo: {prestamo.tercero}",
            usuario=request.user,
        )

        messages.success(request, "Devolución registrada correctamente.")

    return redirect("inventario:prestamos")

@login_required
def productos_inventario(request):
    empresa = obtener_empresa_usuario(request)

    productos = ProductoInventario.objects.filter(
        empresa=empresa
    ).order_by("tipo", "nombre") if empresa else ProductoInventario.objects.none()

    return render(request, "inventario/productos.html", {
        "titulo": "Materia Prima e Inventario",
        "productos": productos,
    })


@login_required
def descargar_plantilla_inventario(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario"

    ws.append([
        "codigo",
        "nombre",
        "tipo",
        "unidad_medida",
        "unidad_compra",
        "cantidad_por_empaque",
        "stock_actual",
        "stock_minimo",
        "precio_unitario_compra",
        "porcentaje_itbis",
        "proveedor",
        "activo",
    ])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="plantilla_inventario.xlsx"'

    wb.save(response)
    return response


@login_required
def cargar_inventario_excel(request):
    empresa = obtener_empresa_usuario(request)

    if request.method == "POST":
        archivo = request.FILES.get("archivo")

        if not archivo:
            messages.error(request, "Debe seleccionar un archivo Excel.")
            return redirect("inventario:productos")

        wb = load_workbook(archivo)
        ws = wb.active

        creados = 0
        actualizados = 0

        for row in ws.iter_rows(min_row=2, values_only=True):
            (
                codigo,
                nombre,
                tipo,
                unidad_medida,
                unidad_compra,
                cantidad_por_empaque,
                stock_actual,
                stock_minimo,
                precio_unitario_compra,
                porcentaje_itbis,
                proveedor,
                activo,
            ) = row

            if not codigo or not nombre:
                continue

            activo_valor = True

            if str(activo).strip().lower() in ["no", "false", "0", "inactivo", "falso"]:
                activo_valor = False

            producto, creado = ProductoInventario.objects.update_or_create(
                empresa=empresa,
                codigo=str(codigo).strip(),
                defaults={
                    "nombre": str(nombre).strip(),
                    "tipo": str(tipo or "materia_prima").strip().lower(),
                    "unidad_medida": str(unidad_medida or "lb").strip().lower(),
                    "unidad_compra": str(unidad_compra or "").strip(),
                    "cantidad_por_empaque": convertir_decimal(cantidad_por_empaque, "1"),
                    "stock_minimo": convertir_decimal(stock_minimo, "0"),
                    "precio_unitario_compra": convertir_decimal(precio_unitario_compra, "0"),
                    "porcentaje_itbis": convertir_itbis(porcentaje_itbis),
                    "proveedor": str(proveedor or "").strip(),
                    "activo": activo_valor,
                }
            )
            saldo_objetivo = convertir_decimal(stock_actual, "0")
            diferencia = saldo_objetivo - Decimal(producto.stock_actual or 0)
            if diferencia:
                _movimiento_desde_vista(
                    request, empresa=empresa, producto=producto,
                    tipo="ajuste" if diferencia > 0 else "salida",
                    cantidad=abs(diferencia),
                    referencia=f"IMPORT-{producto.codigo}-{saldo_objetivo}",
                    observacion="Ajuste autorizado por importacion de inventario.",
                )

            if creado:
                creados += 1
            else:
                actualizados += 1

        messages.success(
            request,
            f"Inventario cargado correctamente. Creados: {creados}. Actualizados: {actualizados}."
        )

        return redirect("inventario:productos")

    return redirect("inventario:productos")


@login_required
def recetas(request):
    empresa = obtener_empresa_usuario(request)

    recetas_qs = Receta.objects.filter(
        empresa=empresa
    ).select_related("producto") if empresa else Receta.objects.none()

    productos_produccion = ProductoProduccion.objects.filter(
        empresa=empresa,
        activo=True
    ).order_by("tipo", "nombre") if empresa else ProductoProduccion.objects.none()

    return render(request, "inventario/recetas.html", {
        "titulo": "Recetas de Producción",
        "recetas": recetas_qs,
        "productos_produccion": productos_produccion,
    })


@login_required
def prestamos(request):
    empresa = obtener_empresa_usuario(request)

    prestamos_qs = PrestamoMateriaPrima.objects.filter(
        empresa=empresa
    ).select_related("producto") if empresa else PrestamoMateriaPrima.objects.none()

    return render(request, "inventario/prestamos.html", {
        "titulo": "Préstamos de Materia Prima",
        "prestamos": prestamos_qs,
    })


@login_required
def ordenes_compra(request):
    empresa = obtener_empresa_usuario(request)

    ordenes = OrdenCompra.objects.filter(
        empresa=empresa
    ).order_by("-fecha", "-id") if empresa else OrdenCompra.objects.none()

    return render(request, "inventario/ordenes_compra.html", {
        "titulo": "Órdenes de Compra",
        "ordenes": ordenes,
    })


@login_required
def generar_orden_compra_sugerida(request):
    empresa = obtener_empresa_usuario(request)

    if request.method == "POST":
        orden = OrdenCompra.objects.create(
            empresa=empresa,
            proveedor=request.POST.get("proveedor"),
            ciudad=request.POST.get("ciudad"),
            fecha_inicio=request.POST.get("fecha_inicio"),
            fecha_fin=request.POST.get("fecha_fin"),
            fecha_requerida=request.POST.get("fecha_requerida") or None,
            observacion=request.POST.get("observacion"),
            creada_por=request.user,
            estado="borrador",
        )

        messages.success(
            request,
            f"Orden de compra {orden.numero} creada correctamente."
        )

        return redirect("inventario:detalle_orden_compra", orden.id)

    return render(request, "inventario/generar_orden_compra.html", {
        "titulo": "Generar orden de compra sugerida",
        "empresa": empresa,
    })


@login_required
def detalle_orden_compra(request, orden_id):
    empresa = obtener_empresa_usuario(request)

    orden = get_object_or_404(
        OrdenCompra,
        id=orden_id,
        empresa=empresa
    )

    return render(request, "inventario/detalle_orden_compra.html", {
        "orden": orden,
        "detalles": orden.detalles.all(),
    })


@login_required
def calcular_orden_compra(request, orden_id):
    empresa = obtener_empresa_usuario(request)

    orden = get_object_or_404(
        OrdenCompra,
        id=orden_id,
        empresa=empresa
    )

    if not orden.fecha_inicio or not orden.fecha_fin:
        messages.error(request, "La orden debe tener fecha inicio y fecha fin.")
        return redirect("inventario:detalle_orden_compra", orden.id)

    orden.detalles.all().delete()

    menus = MenuDiario.objects.filter(
        empresa=empresa,
        fecha__range=[orden.fecha_inicio, orden.fecha_fin]
    ).order_by("fecha")

    if not menus.exists():
        messages.error(request, "No hay menú diario registrado para el rango seleccionado.")
        return redirect("inventario:detalle_orden_compra", orden.id)

    centros = CentroEducativo.objects.filter(empresa=empresa)

    total_raciones = sum(
        Decimal(centro.matricula or 0)
        for centro in centros
    )

    if total_raciones <= 0:
        messages.error(request, "No hay matrícula registrada en los centros educativos.")
        return redirect("inventario:detalle_orden_compra", orden.id)

    necesidades = {}
    menus_sin_receta = []

    for menu in menus:
        producto_menu = (menu.producto or "").strip()
        receta = buscar_receta_por_producto(empresa, producto_menu)

        if not receta:
            menus_sin_receta.append(producto_menu)
            continue

        for detalle in receta.detalles.all():
            materia = detalle.materia_prima

            if not receta.rendimiento_unidades:
                continue

            cantidad_por_racion = Decimal(detalle.cantidad or 0) / Decimal(receta.rendimiento_unidades)
            cantidad_necesaria_dia = cantidad_por_racion * total_raciones

            porcentaje_variacion = Decimal(receta.porcentaje_variacion or 0)

            if porcentaje_variacion > 0:
                cantidad_necesaria_dia = cantidad_necesaria_dia * (
                    Decimal("1") + (porcentaje_variacion / Decimal("100"))
                )

            if materia.id not in necesidades:
                necesidades[materia.id] = {
                    "producto": materia,
                    "cantidad_necesaria": Decimal("0"),
                }

            necesidades[materia.id]["cantidad_necesaria"] += cantidad_necesaria_dia

    for item in necesidades.values():
        producto = item["producto"]
        cantidad_necesaria = item["cantidad_necesaria"]
        cantidad_disponible = Decimal(producto.stock_actual or 0)

        cantidad_faltante = cantidad_necesaria - cantidad_disponible

        if cantidad_faltante < 0:
            cantidad_faltante = Decimal("0")

        cantidad_sugerida = cantidad_empaques_sugeridos(
            cantidad_faltante,
            producto.cantidad_por_empaque
        )

        if cantidad_sugerida <= 0:
            continue

        DetalleOrdenCompra.objects.create(
            orden=orden,
            producto=producto,
            unidad_base=producto.unidad_medida,
            unidad_compra=producto.unidad_compra,
            cantidad_necesaria=cantidad_necesaria,
            cantidad_disponible=cantidad_disponible,
            cantidad_faltante=cantidad_faltante,
            cantidad_sugerida_compra=cantidad_sugerida,
            cantidad_compra=cantidad_sugerida,
            precio_unitario_compra=producto.precio_unitario_compra,
            porcentaje_itbis=producto.porcentaje_itbis,
        )

    orden.recalcular_totales()

    if menus_sin_receta:
        messages.warning(
            request,
            "La orden fue calculada, pero algunos productos del menú no tenían receta registrada: "
            + ", ".join(sorted(set(menus_sin_receta)))
        )
    else:
        messages.success(
            request,
            "Orden de compra calculada correctamente según recetas, menú e inventario disponible."
        )

    return redirect("inventario:detalle_orden_compra", orden.id)


# =====================================================
# PDF ORDEN DE COMPRA
# =====================================================

@login_required
def generar_pdf_orden_compra_response(request, orden_id, descargar=False):
    empresa = obtener_empresa_usuario(request)

    orden = get_object_or_404(
        OrdenCompra,
        id=orden_id,
        empresa=empresa
    )

    detalles = orden.detalles.all()

    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=letter)

    width, height = letter
    margen_x = 55
    y = 745

    meses = {
        1: "enero", 2: "febrero", 3: "marzo", 4: "abril",
        5: "mayo", 6: "junio", 7: "julio", 8: "agosto",
        9: "septiembre", 10: "octubre", 11: "noviembre", 12: "diciembre",
    }

    def fecha_larga(fecha):
        if not fecha:
            return "-"
        return f"{fecha.day} de {meses[fecha.month]} de {fecha.year}"

    def pie_pagina():
        pdf.setFont("Helvetica", 7)
        pdf.drawCentredString(width / 2, 25, f"Página {pdf.getPageNumber()}")

    def encabezado():
        y_local = 745

        if empresa and empresa.logo:
            try:
                with empresa.logo.storage.open(
                    empresa.logo.name,
                    "rb",
                ) as archivo_logo:
                    logo = ImageReader(
                        __import__("io").BytesIO(
                            archivo_logo.read()
                        )
                    )
                pdf.drawImage(
                    logo,
                    width / 2 - 45,
                    704,
                    width=90,
                    height=52,
                    preserveAspectRatio=True,
                    mask="auto"
                )
                y_local = 698
            except Exception:
                pass

        pdf.setFont("Helvetica-Bold", 9)
        pdf.drawCentredString(width / 2, y_local, (empresa.nombre or "").upper())

        y_local -= 12
        pdf.setFont("Helvetica", 7)
        pdf.drawCentredString(width / 2, y_local, empresa.direccion or "")

        y_local -= 10
        pdf.drawCentredString(width / 2, y_local, f"Teléfono: {empresa.telefono or ''}")

        y_local -= 10
        pdf.drawCentredString(width / 2, y_local, f"RNC: {empresa.rnc or ''}")

        return y_local - 35

    def nueva_pagina():
        pie_pagina()
        pdf.showPage()
        return encabezado()

    y = encabezado()

    pdf.setFont("Helvetica-Bold", 13)
    pdf.drawCentredString(width / 2, y, "ORDEN DE COMPRA")

    y -= 25

    pdf.setFont("Helvetica-Bold", 8)
    pdf.drawString(margen_x, y, f"No. Orden: {orden.numero}")

    pdf.drawRightString(
    width - margen_x,
    y,
    orden.ciudad or empresa.ciudad or "Santo Domingo Este"
)

    y -= 11

    pdf.setFont("Helvetica", 8)
    pdf.drawRightString(
    width - margen_x,
    y,
    fecha_larga(orden.fecha)
)

    y -= 18

    pdf.drawString(
    margen_x,
    y,
    f"Fecha requerida: {fecha_larga(orden.fecha_requerida)}"
)

    y -= 16

    pdf.setFont("Helvetica-Bold", 8)
    pdf.drawString(margen_x, y, "Proveedor:")

    pdf.setFont("Helvetica", 8)
    pdf.drawString(
    margen_x + 65,
    y,
    orden.proveedor or "-"
)

    y -= 16

    pdf.setFont("Helvetica-Bold", 8)
    pdf.drawString(margen_x, y, "Período cubierto:")

    pdf.setFont("Helvetica", 8)
    pdf.drawString(
    margen_x + 92,
    y,
    f"{fecha_larga(orden.fecha_inicio)} al {fecha_larga(orden.fecha_fin)}"
)

    y -= 25

    estilo_producto = ParagraphStyle(
        name="ProductoOC",
        fontName="Helvetica",
        fontSize=7,
        leading=8,
    )

    data = [[
        "No.",
        "Producto",
        "Unidad de compra",
        "Cantidad solicitada",
        "Observación",
    ]]

    for index, detalle in enumerate(detalles, start=1):
        data.append([
            str(index),
            Paragraph(detalle.nombre_producto(), estilo_producto),
            detalle.unidad_compra or "-",
            f"{detalle.cantidad_compra:,.2f}",
            Paragraph(detalle.observacion or "", estilo_producto),
        ])

    if not detalles.exists():
        data.append([
            "-",
            "No hay productos registrados en esta orden.",
            "-",
            "-",
            "-",
        ])

    col_widths = [30, 220, 110, 100, 95]

    def crear_tabla(tabla_data):
        tabla = Table(tabla_data, colWidths=col_widths, repeatRows=1)
        tabla.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.35, colors.black),
            ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 6.5),
            ("ALIGN", (0, 0), (-1, 0), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 1), (-1, -1), 7),
            ("ALIGN", (0, 1), (0, -1), "CENTER"),
            ("ALIGN", (2, 1), (3, -1), "CENTER"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        return tabla

    encabezado_tabla = data[0]
    filas_tabla = data[1:]
    indice = 0

    while indice < len(filas_tabla):
        espacio_disponible = y - 95

        if espacio_disponible < 100:
            y = nueva_pagina()
            espacio_disponible = y - 95

        tabla = None
        alto_tabla = 0
        filas_que_caben = 0
        max_filas = len(filas_tabla) - indice

        for cantidad_filas in range(max_filas, 0, -1):
            bloque = filas_tabla[indice:indice + cantidad_filas]
            tabla_prueba = crear_tabla([encabezado_tabla] + bloque)
            _, alto_prueba = tabla_prueba.wrap(0, 0)

            if alto_prueba <= espacio_disponible:
                tabla = tabla_prueba
                alto_tabla = alto_prueba
                filas_que_caben = cantidad_filas
                break

        if tabla is None:
            y = nueva_pagina()
            continue

        tabla.drawOn(pdf, margen_x, y - alto_tabla)
        y = y - alto_tabla - 20
        indice += filas_que_caben

    if orden.observacion:
        if y < 150:
            y = nueva_pagina()

        pdf.setFont("Helvetica-Bold", 8)
        pdf.drawString(margen_x, y, "Observación general:")

        y -= 13

        estilo_obs = ParagraphStyle(
            name="ObsOC",
            fontName="Helvetica",
            fontSize=8,
            leading=11,
        )

        parrafo_obs = Paragraph(orden.observacion, estilo_obs)
        _, alto_obs = parrafo_obs.wrap(width - (margen_x * 2), 70)
        parrafo_obs.drawOn(pdf, margen_x, y - alto_obs)

        y = y - alto_obs - 25

    if y < 150:
        y = nueva_pagina()

    pdf.setFont("Helvetica", 8)
    pdf.drawString(margen_x, y, "Atentamente,")

    y -= 55

    pdf.line(margen_x, y, margen_x + 210, y)
    pdf.line(width - margen_x - 210, y, width - margen_x, y)

    y -= 12

    pdf.setFont("Helvetica-Bold", 8)
    pdf.drawString(margen_x, y, "SOLICITADO POR")
    pdf.drawString(width - margen_x - 210, y, "RECIBIDO / PROVEEDOR")

    pie_pagina()

    pdf.save()
    buffer.seek(0)

    response = HttpResponse(buffer, content_type="application/pdf")

    if descargar:
        response["Content-Disposition"] = f'attachment; filename="orden_compra_{orden.numero}.pdf"'
    else:
        response["Content-Disposition"] = f'inline; filename="orden_compra_{orden.numero}.pdf"'

    return response


@login_required
def pdf_orden_compra(request, orden_id):
    return generar_pdf_orden_compra_response(
        request,
        orden_id,
        descargar=False
    )


@login_required
def descargar_pdf_orden_compra(request, orden_id):
    return generar_pdf_orden_compra_response(
        request,
        orden_id,
        descargar=True
    )

@login_required
def actualizar_detalle_orden(request, detalle_id):
    empresa = obtener_empresa_usuario(request)

    detalle = get_object_or_404(
        DetalleOrdenCompra,
        id=detalle_id,
        orden__empresa=empresa
    )

    if request.method == "POST":

        detalle.cantidad_compra = convertir_decimal(
            request.POST.get("cantidad_compra"),
            "0"
        )

        detalle.precio_unitario_compra = convertir_decimal(
            request.POST.get("precio_unitario_compra"),
            "0"
        )

        detalle.observacion = request.POST.get("observacion", "").strip()

        detalle.save()

        detalle.orden.recalcular_totales()

        messages.success(
            request,
            "Producto actualizado correctamente."
        )

    return redirect(
        "inventario:detalle_orden_compra",
        detalle.orden.id
    )


@login_required
def eliminar_detalle_orden(request, detalle_id):
    empresa = obtener_empresa_usuario(request)

    detalle = get_object_or_404(
        DetalleOrdenCompra,
        id=detalle_id,
        orden__empresa=empresa
    )

    orden = detalle.orden

    detalle.delete()

    orden.recalcular_totales()

    messages.success(
        request,
        "Producto eliminado correctamente."
    )

    return redirect(
        "inventario:detalle_orden_compra",
        orden.id
    )
@login_required
def agregar_producto_manual_orden(request, orden_id):
    empresa = obtener_empresa_usuario(request)

    orden = get_object_or_404(
        OrdenCompra,
        id=orden_id,
        empresa=empresa
    )

    if request.method == "POST":
        DetalleOrdenCompra.objects.create(
            orden=orden,
            producto_manual=request.POST.get("producto_manual", "").strip(),
            unidad_compra=request.POST.get("unidad_compra", "").strip(),
            cantidad_necesaria=convertir_decimal(request.POST.get("cantidad_compra"), "0"),
            cantidad_disponible=Decimal("0"),
            cantidad_faltante=convertir_decimal(request.POST.get("cantidad_compra"), "0"),
            cantidad_sugerida_compra=convertir_decimal(request.POST.get("cantidad_compra"), "0"),
            cantidad_compra=convertir_decimal(request.POST.get("cantidad_compra"), "0"),
            precio_unitario_compra=convertir_decimal(request.POST.get("precio_unitario_compra"), "0"),
            porcentaje_itbis=Decimal("0"),
            observacion=request.POST.get("observacion", "").strip(),
        )

        orden.recalcular_totales()

        messages.success(request, "Producto manual agregado correctamente.")

    return redirect("inventario:detalle_orden_compra", orden.id)
# =====================================================
# RECIBIR ORDEN DE COMPRA
# =====================================================

@login_required
def recibir_orden_compra(request, orden_id):
    empresa = obtener_empresa_usuario(request)

    orden = get_object_or_404(
        OrdenCompra,
        id=orden_id,
        empresa=empresa
    )

    if orden.inventario_actualizado:
        messages.warning(
            request,
            "Esta orden ya fue recibida anteriormente."
        )

        return redirect(
            "inventario:detalle_orden_compra",
            orden.id
        )

    detalles = orden.detalles.all()

    if not detalles.exists():
        messages.error(
            request,
            "La orden no tiene productos."
        )

        return redirect(
            "inventario:detalle_orden_compra",
            orden.id
        )

    for detalle in detalles:

        # Solo productos reales de inventario
        if not detalle.producto:
            continue

        producto = detalle.producto

        cantidad_empaques = Decimal(detalle.cantidad_compra or 0)

        cantidad_por_empaque = Decimal(
            producto.cantidad_por_empaque or 1
        )

        cantidad_total_base = (
            cantidad_empaques * cantidad_por_empaque
        )

        _movimiento_desde_vista(request,
            empresa=empresa,
            producto=producto,
            tipo="entrada_compra",
            cantidad=cantidad_total_base,
            costo_unitario=detalle.precio_unitario_compra or 0,
            fecha=timezone.localdate(),
            referencia=orden.numero,
            observacion=f"Entrada automática por orden de compra {orden.numero}",
            usuario=request.user,
        )

    orden.estado = "recibida"
    orden.fecha_recepcion = timezone.localdate()
    orden.recibida_por = request.user
    orden.inventario_actualizado = True

    orden.save(
        update_fields=[
            "estado",
            "fecha_recepcion",
            "recibida_por",
            "inventario_actualizado",
        ]
    )

    messages.success(
        request,
        f"Orden {orden.numero} recibida e inventario actualizado correctamente."
    )

    return redirect(
        "inventario:detalle_orden_compra",
        orden.id
    )
# =====================================================
# PRODUCCIÓN AUTOMÁTICA
# =====================================================

@login_required
def produccion(request):
    empresa = obtener_empresa_usuario(request)

    producciones = ProduccionProgramada.objects.filter(
        empresa=empresa
    ).select_related(
        "producto",
        "receta",
        "usuario"
    ).order_by("-fecha", "-id") if empresa else ProduccionProgramada.objects.none()

    return render(request, "inventario/produccion.html", {
        "producciones": producciones,
    })


@login_required
def generar_produccion_desde_menu(request):
    empresa = obtener_empresa_usuario(request)

    if request.method == "POST":
        fecha_inicio = request.POST.get("fecha_inicio")
        fecha_fin = request.POST.get("fecha_fin")

        if not fecha_inicio or not fecha_fin:
            messages.error(request, "Debe seleccionar fecha inicio y fecha fin.")
            return redirect("inventario:produccion")

        menus = MenuDiario.objects.filter(
            empresa=empresa,
            fecha__range=[fecha_inicio, fecha_fin]
        ).order_by("fecha")

        if not menus.exists():
            messages.error(request, "No hay menú diario registrado para ese rango.")
            return redirect("inventario:produccion")

        centros = CentroEducativo.objects.filter(empresa=empresa)

        total_raciones = sum(
            Decimal(centro.matricula or 0)
            for centro in centros
        )

        if total_raciones <= 0:
            messages.error(request, "No hay matrícula registrada en los centros educativos.")
            return redirect("inventario:produccion")

        creadas = 0
        sin_receta = []

        for menu in menus:
            producto_menu = (menu.producto or "").strip()
            receta = buscar_receta_por_producto(empresa, producto_menu)

            if not receta:
                sin_receta.append(producto_menu)
                continue

            ProduccionProgramada.objects.get_or_create(
                empresa=empresa,
                fecha=menu.fecha,
                producto=receta.producto,
                receta=receta,
                defaults={
                    "cantidad_unidades": total_raciones,
                    "estado": "programada",
                    "usuario": request.user,
                    "observacion": f"Producción generada automáticamente desde menú: {producto_menu}",
                }
            )

            creadas += 1

        if creadas:
            messages.success(request, f"Producciones generadas correctamente: {creadas}.")

        if sin_receta:
            messages.warning(
                request,
                "Algunos productos del menú no tienen receta registrada: "
                + ", ".join(sorted(set(sin_receta)))
            )

        return redirect("inventario:produccion")

    return redirect("inventario:produccion")


@login_required
def detalle_produccion(request, produccion_id):
    empresa = obtener_empresa_usuario(request)

    produccion = get_object_or_404(
        ProduccionProgramada,
        id=produccion_id,
        empresa=empresa
    )

    detalles = []

    for detalle in produccion.receta.detalles.all():
        cantidad_por_unidad = Decimal(detalle.cantidad or 0) / Decimal(produccion.receta.rendimiento_unidades or 1)
        cantidad_necesaria = cantidad_por_unidad * Decimal(produccion.cantidad_unidades or 0)

        porcentaje_variacion = Decimal(produccion.receta.porcentaje_variacion or 0)

        if porcentaje_variacion > 0:
            cantidad_necesaria = cantidad_necesaria * (
                Decimal("1") + (porcentaje_variacion / Decimal("100"))
            )

        detalles.append({
            "producto": detalle.materia_prima,
            "cantidad_necesaria": cantidad_necesaria,
            "stock_actual": detalle.materia_prima.stock_actual,
            "unidad": detalle.materia_prima.get_unidad_medida_display(),
            "costo_estimado": cantidad_necesaria * detalle.materia_prima.costo_unitario,
        })

    costo_estimado = sum(item["costo_estimado"] for item in detalles)

    return render(request, "inventario/detalle_produccion.html", {
        "produccion": produccion,
        "detalles": detalles,
        "costo_estimado": costo_estimado,
    })


@login_required
def ejecutar_produccion(request, produccion_id):
    empresa = obtener_empresa_usuario(request)

    produccion = get_object_or_404(
        ProduccionProgramada,
        id=produccion_id,
        empresa=empresa
    )

    if produccion.estado == "ejecutada":
        messages.warning(request, "Esta producción ya fue ejecutada anteriormente.")
        return redirect("inventario:detalle_produccion", produccion.id)

    if not produccion.receta or not produccion.receta.detalles.exists():
        messages.error(request, "Esta producción no tiene receta o ingredientes registrados.")
        return redirect("inventario:detalle_produccion", produccion.id)

    costo_real = Decimal("0")

    for detalle in produccion.receta.detalles.all():
        materia = detalle.materia_prima

        cantidad_por_unidad = Decimal(detalle.cantidad or 0) / Decimal(produccion.receta.rendimiento_unidades or 1)
        cantidad_consumir = cantidad_por_unidad * Decimal(produccion.cantidad_unidades or 0)

        porcentaje_variacion = Decimal(produccion.receta.porcentaje_variacion or 0)

        if porcentaje_variacion > 0:
            cantidad_consumir = cantidad_consumir * (
                Decimal("1") + (porcentaje_variacion / Decimal("100"))
            )

        costo_real += cantidad_consumir * materia.costo_unitario

        _movimiento_desde_vista(request,
            empresa=empresa,
            producto=materia,
            tipo="produccion",
            cantidad=cantidad_consumir,
            costo_unitario=materia.costo_unitario,
            fecha=produccion.fecha,
            referencia=f"PROD-{produccion.id}",
            observacion=f"Consumo automático por producción de {produccion.producto.nombre}",
            usuario=request.user,
        )

    produccion.estado = "ejecutada"
    produccion.ejecutada_en = timezone.now()
    produccion.usuario = request.user

    # Si agregaste estos campos al modelo, se guardarán; si no, omite estas dos líneas.
    if hasattr(produccion, "costo_real"):
        produccion.costo_real = costo_real

    produccion.save()

    messages.success(
        request,
        f"Producción ejecutada correctamente. Costo estimado: RD$ {costo_real:,.2f}"
    )

    return redirect("inventario:detalle_produccion", produccion.id)
@login_required
def generar_proyeccion_consumo(request):

    empresa = obtener_empresa_usuario(request)

    fecha_inicio = request.POST.get("fecha_inicio")
    fecha_fin = request.POST.get("fecha_fin")

    menus = MenuDiario.objects.filter(
        empresa=empresa,
        fecha__range=[fecha_inicio, fecha_fin]
    )

    centros = CentroEducativo.objects.filter(
        empresa=empresa
    )

    total_raciones = sum(
        Decimal(c.matricula or 0)
        for c in centros
    )

    ProyeccionConsumoDiario.objects.filter(
        empresa=empresa,
        fecha__range=[fecha_inicio, fecha_fin]
    ).delete()

    for menu in menus:

        receta = buscar_receta_por_producto(
            empresa,
            menu.producto
        )

        if not receta:
            continue

        for detalle in receta.detalles.all():

            cantidad_por_racion = (
                Decimal(detalle.cantidad)
                / Decimal(receta.rendimiento_unidades)
            )

            cantidad_total = (
                cantidad_por_racion
                * total_raciones
            )

            if receta.porcentaje_variacion:
                cantidad_total *= (
                    Decimal("1")
                    + (
                        Decimal(receta.porcentaje_variacion)
                        / Decimal("100")
                    )
                )

            ProyeccionConsumoDiario.objects.create(
                empresa=empresa,
                fecha=menu.fecha,
                producto_menu=menu.producto,
                receta=receta,
                materia_prima=detalle.materia_prima,
                raciones=total_raciones,
                cantidad_proyectada=cantidad_total,
                generado_por=request.user
            )

    messages.success(
        request,
        "Proyección automática generada correctamente."
    )

    return redirect("inventario:dashboard")

@login_required
def registrar_consumo_manual(request, produccion_id):

    empresa = obtener_empresa_usuario(request)

    produccion = get_object_or_404(
        ProduccionProgramada,
        id=produccion_id,
        empresa=empresa
    )

    if request.method == "POST":

        producto_id = request.POST.get("producto")
        cantidad = convertir_decimal(
            request.POST.get("cantidad")
        )

        observacion = request.POST.get("observacion")

        producto = get_object_or_404(
            ProductoInventario,
            id=producto_id,
            empresa=empresa
        )

        DetalleConsumoProduccion.objects.create(
            produccion=produccion,
            producto=producto,
            cantidad=cantidad,
            observacion=observacion,
            usuario=request.user
        )

        _movimiento_desde_vista(request,
            empresa=empresa,
            producto=producto,
            tipo="salida",
            cantidad=cantidad,
            observacion=f"Consumo manual producción #{produccion.id}",
            usuario=request.user
        )

        messages.success(
            request,
            "Consumo manual registrado."
        )

    return redirect(
        "inventario:detalle_produccion",
        produccion.id
    )
@login_required
def pdf_inventario(request):

    empresa = obtener_empresa_usuario(request)

    productos = ProductoInventario.objects.filter(
        empresa=empresa,
        activo=True
    ).order_by("nombre")

    response = HttpResponse(
        content_type="application/pdf"
    )

    response["Content-Disposition"] = (
        'inline; filename="inventario.pdf"'
    )

    pdf = canvas.Canvas(response, pagesize=letter)

    width, height = letter

    y = height - 50

    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(40, y, "Inventario disponible")

    y -= 30

    data = [[
        "Producto",
        "Tipo",
        "Unidad",
        "Disponible",
        "Mínimo",
        "Estado"
    ]]

    for producto in productos:

        estado = "OK"

        if producto.esta_bajo_minimo():
            estado = "CRÍTICO"

        data.append([
            producto.nombre,
            producto.get_tipo_display(),
            producto.get_unidad_medida_display(),
            f"{producto.stock_actual}",
            f"{producto.stock_minimo}",
            estado,
        ])

    tabla = Table(data, colWidths=[
        180,
        90,
        70,
        70,
        70,
        70,
    ])

    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
    ]))

    tabla.wrapOn(pdf, width, height)
    tabla.drawOn(pdf, 40, y - (len(data) * 18))

    pdf.save()

    return response

# =====================================================
# CRUD RECETAS
# =====================================================

@login_required
def crear_producto_produccion(request):
    empresa = obtener_empresa_usuario(request)

    if request.method == "POST":
        ProductoProduccion.objects.create(
            empresa=empresa,
            nombre=request.POST.get("nombre", "").strip(),
            tipo=request.POST.get("tipo"),
            unidad_resultado="unidad",
            activo=True,
        )

        messages.success(request, "Producto de producción creado correctamente.")

    return redirect("inventario:recetas")


@login_required
def crear_receta(request):
    empresa = obtener_empresa_usuario(request)

    if request.method == "POST":
        producto = get_object_or_404(
            ProductoProduccion,
            id=request.POST.get("producto_id"),
            empresa=empresa
        )

        Receta.objects.create(
            empresa=empresa,
            producto=producto,
            nombre=request.POST.get("nombre", "").strip(),
            rendimiento_unidades=convertir_decimal(request.POST.get("rendimiento_unidades"), "1"),
            porcentaje_variacion=convertir_decimal(request.POST.get("porcentaje_variacion"), "0"),
            observaciones=request.POST.get("observaciones", "").strip(),
            activa=True,
        )

        messages.success(request, "Receta creada correctamente.")

    return redirect("inventario:recetas")


@login_required
def detalle_receta(request, receta_id):
    empresa = obtener_empresa_usuario(request)

    receta = get_object_or_404(
        Receta,
        id=receta_id,
        empresa=empresa
    )

    productos_inventario = ProductoInventario.objects.filter(
        empresa=empresa,
        activo=True
    ).order_by("tipo", "nombre")

    detalles = receta.detalles.select_related("materia_prima")

    costo_total = Decimal("0")

    for detalle in detalles:
        costo_total += Decimal(detalle.cantidad or 0) * Decimal(detalle.materia_prima.costo_unitario or 0)

    costo_por_unidad = Decimal("0")

    if receta.rendimiento_unidades:
        costo_por_unidad = costo_total / Decimal(receta.rendimiento_unidades)

    return render(request, "inventario/detalle_receta.html", {
        "receta": receta,
        "detalles": detalles,
        "productos_inventario": productos_inventario,
        "costo_total": costo_total,
        "costo_por_unidad": costo_por_unidad,
    })


@login_required
def agregar_ingrediente_receta(request, receta_id):
    empresa = obtener_empresa_usuario(request)

    receta = get_object_or_404(
        Receta,
        id=receta_id,
        empresa=empresa
    )

    if request.method == "POST":
        materia_prima = get_object_or_404(
            ProductoInventario,
            id=request.POST.get("materia_prima_id"),
            empresa=empresa
        )

        DetalleReceta.objects.update_or_create(
            receta=receta,
            materia_prima=materia_prima,
            defaults={
                "cantidad": convertir_decimal(request.POST.get("cantidad"), "0"),
                "unidad_medida": request.POST.get("unidad_medida") or materia_prima.unidad_medida,
                "porcentaje": convertir_decimal(request.POST.get("porcentaje"), "0"),
            }
        )

        messages.success(request, "Ingrediente agregado correctamente.")

    return redirect("inventario:detalle_receta", receta.id)


@login_required
def eliminar_ingrediente_receta(request, detalle_id):
    empresa = obtener_empresa_usuario(request)

    detalle = get_object_or_404(
        DetalleReceta,
        id=detalle_id,
        receta__empresa=empresa
    )

    receta_id = detalle.receta.id
    detalle.delete()

    messages.success(request, "Ingrediente eliminado correctamente.")

    return redirect("inventario:detalle_receta", receta_id)
# =====================================================
# MOVIMIENTOS
# =====================================================

@login_required
def movimientos(request):
    empresa = obtener_empresa_usuario(request)

    movimientos_qs = (
        MovimientoInventario.objects
        .filter(empresa=empresa)
        .select_related("producto", "usuario")
        .order_by("-fecha", "-id")
        if empresa
        else MovimientoInventario.objects.none()
    )

    contexto = {
        "titulo": "Movimientos de Inventario",
        "movimientos": movimientos_qs,
    }

    return render(
        request,
        "inventario/movimientos.html",
        contexto
    )
# =====================================================
# EDICIÓN DE INVENTARIO / KARDEX / MOVIMIENTOS MANUALES
# =====================================================

@login_required
def editar_producto_inventario(request, producto_id):
    empresa = obtener_empresa_usuario(request)

    producto = get_object_or_404(
        ProductoInventario,
        id=producto_id,
        empresa=empresa
    )

    if request.method == "POST":
        producto.codigo = request.POST.get("codigo", "").strip()
        producto.nombre = request.POST.get("nombre", "").strip()
        producto.tipo = request.POST.get("tipo")
        producto.unidad_medida = request.POST.get("unidad_medida")
        producto.unidad_compra = request.POST.get("unidad_compra", "").strip()
        producto.cantidad_por_empaque = convertir_decimal(request.POST.get("cantidad_por_empaque"), "1")
        producto.stock_minimo = convertir_decimal(request.POST.get("stock_minimo"), "0")
        producto.precio_unitario_compra = convertir_decimal(request.POST.get("precio_unitario_compra"), "0")
        producto.porcentaje_itbis = convertir_itbis(request.POST.get("porcentaje_itbis"))
        producto.proveedor = request.POST.get("proveedor", "").strip()
        producto.activo = request.POST.get("activo") == "on"
        producto.clasificacion_operativa = request.POST.get(
    "clasificacion_operativa"
)

        producto.afecta_produccion = (
    request.POST.get("afecta_produccion") == "on"
)
        producto.save()

        messages.success(request, "Producto actualizado correctamente.")
        return redirect("inventario:productos")

    return render(request, "inventario/editar_producto.html", {
        "producto": producto,
    })


@login_required
def desactivar_producto_inventario(request, producto_id):
    empresa = obtener_empresa_usuario(request)

    producto = get_object_or_404(
        ProductoInventario,
        id=producto_id,
        empresa=empresa
    )

    producto.activo = False
    producto.save(update_fields=["activo"])

    messages.success(request, "Producto desactivado correctamente.")
    return redirect("inventario:productos")

@login_required
def generar_orden_compra_sugerida(request):
    empresa = obtener_empresa_usuario(request)

    if request.method == "POST":
        fecha_inicio = request.POST.get("fecha_inicio")
        fecha_fin = request.POST.get("fecha_fin")
        proveedor = request.POST.get("proveedor")
        observacion = request.POST.get("observacion")

        orden = OrdenCompra.objects.create(
            empresa=empresa,
            proveedor=proveedor,
            fecha_inicio=fecha_inicio,
            fecha_fin=fecha_fin,
            observacion=observacion,
            creada_por=request.user,
            estado="borrador",
        )

        messages.success(
            request,
            f"Orden de compra {orden.numero} creada correctamente."
        )

        return redirect("inventario:detalle_orden_compra", orden.id)

    contexto = {
        "titulo": "Generar orden de compra sugerida",
        "empresa": empresa,
    }

    return render(request, "inventario/generar_orden_compra.html", contexto)


@login_required
def detalle_orden_compra(request, orden_id):
    empresa = obtener_empresa_usuario(request)

    orden = get_object_or_404(
        OrdenCompra,
        id=orden_id,
        empresa=empresa
    )

    detalles = orden.detalles.all()

    contexto = {
        "orden": orden,
        "detalles": detalles,
    }

    return render(request, "inventario/detalle_orden_compra.html", contexto)


@login_required
def registrar_movimiento_manual(request):
    empresa = obtener_empresa_usuario(request)

    productos = ProductoInventario.objects.filter(
        empresa=empresa,
        activo=True
    ).order_by("nombre")

    if request.method == "POST":
        producto = get_object_or_404(
            ProductoInventario,
            id=request.POST.get("producto_id"),
            empresa=empresa
        )

        tipo = request.POST.get("tipo")
        cantidad = convertir_decimal(request.POST.get("cantidad"), "0")
        costo_unitario = convertir_decimal(request.POST.get("costo_unitario"), "0")
        referencia = request.POST.get("referencia", "").strip()
        observacion = request.POST.get("observacion", "").strip()

        if cantidad <= 0:
            messages.error(request, "La cantidad debe ser mayor que cero.")
            return redirect("inventario:registrar_movimiento_manual")

        _movimiento_desde_vista(request,
            empresa=empresa,
            producto=producto,
            tipo=tipo,
            cantidad=cantidad,
            costo_unitario=costo_unitario,
            fecha=timezone.localdate(),
            referencia=referencia,
            observacion=observacion,
            usuario=request.user,
        )

        messages.success(request, "Movimiento registrado correctamente.")
        return redirect("inventario:movimientos")

    return render(request, "inventario/registrar_movimiento.html", {
        "productos": productos,
    })


@login_required
def kardex_producto(request, producto_id):
    empresa = obtener_empresa_usuario(request)

    producto = get_object_or_404(
        ProductoInventario,
        id=producto_id,
        empresa=empresa
    )

    movimientos = MovimientoInventario.objects.filter(
        empresa=empresa,
        producto=producto
    ).select_related("usuario").order_by("fecha", "id")

    balance = Decimal("0")
    filas = []

    for movimiento in movimientos:
        entrada = Decimal("0")
        salida = Decimal("0")

        if movimiento.tipo in [
            "entrada",
            "entrada_compra",
            "ajuste",
            "prestamo_recibido",
            "devolucion_prestamo",
        ]:
            entrada = Decimal(movimiento.cantidad or 0)
            balance += entrada
        else:
            salida = Decimal(movimiento.cantidad or 0)
            balance -= salida

        filas.append({
            "movimiento": movimiento,
            "entrada": entrada,
            "salida": salida,
            "balance": balance,
        })

    return render(request, "inventario/kardex_producto.html", {
        "producto": producto,
        "filas": filas,
    })


@login_required
def eliminar_orden_compra(request, orden_id):
    empresa = obtener_empresa_usuario(request)

    orden = get_object_or_404(
        OrdenCompra,
        id=orden_id,
        empresa=empresa
    )

    if orden.inventario_actualizado:
        messages.error(
            request,
            "No se puede eliminar una orden ya recibida porque actualizó el inventario."
        )
        return redirect("inventario:detalle_orden_compra", orden.id)

    orden.delete()

    messages.success(request, "Orden de compra eliminada correctamente.")
    return redirect("inventario:ordenes_compra")


for _legacy_mutation in (
    "generar_orden_compra_sugerida", "calcular_orden_compra",
    "actualizar_detalle_orden", "eliminar_detalle_orden",
    "agregar_producto_manual_orden", "eliminar_orden_compra",
):
    globals()[_legacy_mutation] = _bloquear_mutacion_orden_legada(globals()[_legacy_mutation])
